from __future__ import annotations

import argparse
import csv
import concurrent.futures
import html
import json
import os
import re
import shutil
import subprocess
import textwrap
import time
import zipfile
from dataclasses import dataclass
from datetime import datetime, timedelta
from pathlib import Path
from typing import Dict, List, Optional, Set, Tuple
from urllib.parse import quote_plus, urljoin
from urllib.parse import urlparse

import httpx
import yaml
from bs4 import BeautifulSoup
from dotenv import load_dotenv
from openpyxl import Workbook
from openpyxl import load_workbook
from openpyxl.chart import BarChart
from openpyxl.chart import LineChart
from openpyxl.chart import PieChart
from openpyxl.chart import Reference
from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pipeline_orchestrator import optimize_application
try:
    from playwright.sync_api import TimeoutError as PlaywrightTimeoutError
    from playwright.sync_api import sync_playwright
    from playwright.sync_api import Error as PlaywrightError
    PLAYWRIGHT_AVAILABLE = True
except Exception:
    PlaywrightTimeoutError = TimeoutError  # type: ignore[assignment]
    sync_playwright = None  # type: ignore[assignment]
    PlaywrightError = Exception  # type: ignore[assignment]
    PLAYWRIGHT_AVAILABLE = False


ROOT = Path(__file__).resolve().parents[1]
SETTINGS_FILE = ROOT / "config" / "settings.yaml"
PROFILE_FILE = ROOT / "config" / "candidate_profile.json"
JOBS_FILE = ROOT / "jobs" / "jobs.csv"
OUTPUT_APPLICATION_DIR = ROOT / "outputs" / "applications"
PROFILE_DIR = ROOT / "browser_profiles"
AUDIT_XLSX_FILE = ROOT / "application_audit.xlsx"
TRACKER_CSV_FILE = ROOT / "outputs" / "reports" / "application_tracker.csv"
TRACKER_MD_FILE = ROOT / "outputs" / "reports" / "application_tracker.md"
EXTERNAL_FACTS_FILE = ROOT / "config" / "external_profile_facts.json"
PORTALS_FILE = ROOT / "config" / "portals.yaml"
SEEK_STRATEGY_FILE = ROOT / "config" / "seek_search_strategy.yaml"
SCAN_HISTORY_FILE = ROOT / "outputs" / "reports" / "scan_history.tsv"
PIPELINE_FILE = ROOT / "outputs" / "reports" / "pipeline.md"
SEEK_STRATEGY_REPORT_FILE = ROOT / "outputs" / "reports" / "seek_search_strategy.md"
PLAYWRIGHT_BROWSERS_DIR = ROOT / ".pw-browsers"


AUDIT_FOCUS_HEADERS = [
    "pick_id",
    "status",
    "match_score",
    "posted_priority",
    "job_posted_date",
    "days_since_posted",
    "company",
    "title",
    "role_lock",
    "platform",
    "location",
    "resume_file",
    "cover_letter_file",
    "matched_tier1_keywords",
    "missing_tier1_keywords",
    "critical_gaps",
    "suggested_improvements",
    "application_url",
    "job_url",
]

AUDIT_DETAIL_HEADERS = [
    "job_id",
    "template_used",
    "keywords_injected",
    "ats_score",
    "section_score",
    "covered_keywords",
    "missing_keywords",
    "improvements",
    "risk_flags",
    "timestamp",
    "applied_at",
    "auto_step",
]

AUDIT_ALWAYS_KEEP_HEADERS = {
    "pick_id",
    "status",
    "match_score",
    "posted_priority",
    "company",
    "title",
    "platform",
    "application_url",
    "job_url",
}


@dataclass
class JobRecord:
    job_id: str
    platform: str
    company: str
    title: str
    location: str
    job_url: str
    application_url: str
    job_posted_date: str
    jd_text: str
    notes: str


CANONICAL_STATUS_ALIASES = {
    "readytoapply": "ReadyToApply",
    "readyforreview": "ReadyForReview",
    "hold": "Hold",
    "submitted": "Submitted",
    "applied": "Submitted",
    "responded": "Responded",
    "interview": "Interview",
    "offer": "Offer",
    "rejected": "Rejected",
    "failed": "Failed",
    "skipped": "Skipped",
    "discarded": "Skipped",
    "skip": "Skipped",
    "skippednojd": "SkippedNoJD",
    "skippednourl": "SkippedNoURL",
}

STATUS_RANK = {
    "Skipped": 0,
    "SkippedNoJD": 0,
    "SkippedNoURL": 0,
    "Failed": 1,
    "Hold": 2,
    "ReadyForReview": 3,
    "ReadyToApply": 4,
    "Submitted": 5,
    "Responded": 6,
    "Interview": 7,
    "Offer": 8,
    "Rejected": 3,
}


def normalize_status_value(raw: str) -> str:
    value = (raw or "").strip()
    if not value:
        return "Hold"
    compact = re.sub(r"[^A-Za-z]", "", value).lower()
    return CANONICAL_STATUS_ALIASES.get(compact, value)


def normalize_posted_date(raw: str) -> str:
    value = str(raw or "").strip()
    if not value:
        return ""
    value = value.replace("Z", "+00:00")
    # Try ISO date/datetime first.
    try:
        dt = datetime.fromisoformat(value)
        return dt.date().isoformat()
    except Exception:
        pass
    # Try extracting YYYY-MM-DD from noisy strings.
    m = re.search(r"(\d{4}-\d{2}-\d{2})", value)
    if m:
        return m.group(1)
    # Fallback for date-only text that may parse.
    for fmt in ["%Y/%m/%d", "%d/%m/%Y", "%m/%d/%Y", "%Y-%m-%d"]:
        try:
            dt = datetime.strptime(value[:10], fmt)
            return dt.date().isoformat()
        except Exception:
            continue
    return ""


def infer_posted_date_from_text(text: str) -> str:
    raw = str(text or "")
    low = raw.lower()
    if not low:
        return ""
    today = datetime.now().date()
    # Absolute date formats.
    for pattern in [
        r"\b(\d{4}-\d{2}-\d{2})\b",
        r"\b(\d{1,2}/\d{1,2}/\d{4})\b",
        r"\b(\d{1,2}\s+[a-z]{3,9}\s+\d{4})\b",
        r"\b([a-z]{3,9}\s+\d{1,2},\s*\d{4})\b",
    ]:
        m_abs = re.search(pattern, low, flags=re.I)
        if not m_abs:
            continue
        parsed = normalize_posted_date(m_abs.group(1))
        if parsed:
            return parsed
    if re.search(r"\b(today|just posted)\b", low):
        return today.isoformat()
    if "yesterday" in low:
        return (today - timedelta(days=1)).isoformat()

    # LinkedIn/Seek style relative formats: "3d ago", "2w ago", "30+ days ago", "reposted 5 days ago".
    m = re.search(r"(?:reposted\s+)?(\d+)\+?\s*(day|days|d|week|weeks|w|month|months|mo)\s+ago", low)
    if m:
        n = int(m.group(1))
        unit = m.group(2)
        if unit in {"d", "day", "days"}:
            delta = timedelta(days=n)
        elif unit in {"w", "week", "weeks"}:
            delta = timedelta(days=n * 7)
        else:
            delta = timedelta(days=n * 30)
        return (today - delta).isoformat()
    m = re.search(r"(?:reposted\s+)?(\d+)\s*(hour|hours|hr|hrs|h|min|mins|minute|minutes)\s+ago", low)
    if m:
        # Sub-day postings are treated as "today" for freshness.
        return today.isoformat()
    return ""


def days_since_posted(posted_date: str) -> str:
    d = normalize_posted_date(posted_date)
    if not d:
        return ""
    try:
        delta = datetime.now().date() - datetime.fromisoformat(d).date()
        return str(max(delta.days, 0))
    except Exception:
        return ""


def posted_priority_bucket(posted_date: str) -> str:
    days = days_since_posted(posted_date)
    if not days.isdigit():
        return "Unknown"
    value = int(days)
    if value <= 2:
        return "Hot"
    if value <= 7:
        return "Fresh"
    if value <= 14:
        return "Aging"
    return "Old"


def parse_numeric_score(raw: str) -> float:
    m = re.search(r"(\d+(?:\.\d+)?)", str(raw or ""))
    return float(m.group(1)) if m else 0.0


def normalize_text_key(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(value or "").lower()).strip()


def fuzzy_title_key(value: str) -> str:
    words = [w for w in normalize_text_key(value).split() if len(w) > 2]
    return " ".join(words[:8])


def build_job_identity(row: dict) -> str:
    url = str(row.get("job_url", "")).strip().lower()
    if url:
        return f"url::{url}"
    company = normalize_text_key(row.get("company", ""))
    title = fuzzy_title_key(row.get("title", ""))
    platform = normalize_text_key(row.get("platform", ""))
    return f"role::{platform}::{company}::{title}"


def merge_row_values(base: dict, incoming: dict) -> dict:
    merged = dict(base)
    for key, value in incoming.items():
        if str(value or "").strip():
            merged[key] = value

    base_status = normalize_status_value(base.get("status", ""))
    incoming_status = normalize_status_value(incoming.get("status", ""))
    if STATUS_RANK.get(incoming_status, 0) >= STATUS_RANK.get(base_status, 0):
        merged["status"] = incoming_status
    else:
        merged["status"] = base_status

    if parse_numeric_score(incoming.get("match_score", "")) > parse_numeric_score(base.get("match_score", "")):
        for field in [
            "resume_file",
            "resume_safe_file",
            "resume_aggressive_file",
            "resume_human_file",
            "cover_letter_file",
            "template_used",
            "role_lock",
            "keywords_injected",
            "match_score",
            "ats_match_score_llm",
            "ats_score",
            "section_score",
            "ats_keywords_used",
            "covered_keywords",
            "missing_keywords",
            "missing_keywords_llm",
            "risk_flags",
            "improvements",
            "remove_for_clarity",
            "strict_profile_facts_used",
            "timestamp",
        ]:
            if str(incoming.get(field, "")).strip():
                merged[field] = incoming.get(field, "")

    if not str(merged.get("job_id", "")).strip():
        merged["job_id"] = incoming.get("job_id") or base.get("job_id") or ""
    return merged


def dedupe_and_normalize_rows(rows: List[dict]) -> List[dict]:
    merged_rows: Dict[str, dict] = {}
    order: List[str] = []
    for row in rows:
        normalized = dict(row)
        normalized["status"] = normalize_status_value(normalized.get("status", ""))
        key = build_job_identity(normalized)
        if key not in merged_rows:
            merged_rows[key] = normalized
            order.append(key)
            continue
        merged_rows[key] = merge_row_values(merged_rows[key], normalized)

    deduped = [merged_rows[key] for key in order]
    deduped.sort(
        key=lambda r: (
            -STATUS_RANK.get(normalize_status_value(r.get("status", "")), 0),
            -parse_numeric_score(r.get("match_score", "")),
            str(r.get("timestamp", "")),
        )
    )
    return deduped


def summarize_tracker(rows: List[dict]) -> dict:
    counts: Dict[str, int] = {}
    score_values: List[float] = []
    with_resume = 0
    with_cover = 0
    high_match = 0
    submitted_or_better = 0
    interview_plus = 0
    fresh_jobs_7d = 0
    aged_jobs_14d = 0
    posted_date_known = 0
    posted_days_values: List[int] = []
    for row in rows:
        status = normalize_status_value(row.get("status", ""))
        counts[status] = counts.get(status, 0) + 1
        score = parse_numeric_score(row.get("match_score", ""))
        if score:
            score_values.append(score)
        if score >= 70:
            high_match += 1
        if str(row.get("resume_file", "")).strip():
            with_resume += 1
        if str(row.get("cover_letter_file", "")).strip():
            with_cover += 1
        if status in {"Submitted", "Responded", "Interview", "Offer"}:
            submitted_or_better += 1
        if status in {"Interview", "Offer"}:
            interview_plus += 1
        posted = normalize_posted_date(row.get("job_posted_date", ""))
        if posted:
            posted_date_known += 1
            days = days_since_posted(posted)
            if days and days.isdigit():
                d = int(days)
                posted_days_values.append(d)
                if d <= 7:
                    fresh_jobs_7d += 1
                if d > 14:
                    aged_jobs_14d += 1
    total = len(rows)
    return {
        "total": total,
        "avg_match_score": round(sum(score_values) / len(score_values), 2) if score_values else 0.0,
        "with_resume_pct": round((with_resume / total) * 100, 1) if total else 0.0,
        "with_cover_pct": round((with_cover / total) * 100, 1) if total else 0.0,
        "high_match_rate_pct": round((high_match / total) * 100, 1) if total else 0.0,
        "submission_rate_pct": round((submitted_or_better / total) * 100, 1) if total else 0.0,
        "interview_plus_rate_pct": round((interview_plus / total) * 100, 1) if total else 0.0,
        "fresh_jobs_7d_pct": round((fresh_jobs_7d / total) * 100, 1) if total else 0.0,
        "aged_jobs_14d_pct": round((aged_jobs_14d / total) * 100, 1) if total else 0.0,
        "posted_date_coverage_pct": round((posted_date_known / total) * 100, 1) if total else 0.0,
        "avg_days_since_posted": round(sum(posted_days_values) / len(posted_days_values), 1) if posted_days_values else 0.0,
        "counts": dict(sorted(counts.items(), key=lambda item: (-STATUS_RANK.get(item[0], 0), item[0]))),
    }


def build_tracker_markdown(rows: List[dict]) -> str:
    export_rows = prepare_export_rows(rows)
    summary = summarize_tracker(export_rows)
    lines = [
        "# Application Tracker",
        "",
        f"- Total roles: {summary['total']}",
        f"- Average match score: {summary['avg_match_score']}",
        f"- Resume coverage: {summary['with_resume_pct']}%",
        f"- Cover letter coverage: {summary['with_cover_pct']}%",
        f"- High-match rate (>=70): {summary.get('high_match_rate_pct', 0.0)}%",
        f"- Submission rate: {summary.get('submission_rate_pct', 0.0)}%",
        f"- Interview-or-better rate: {summary.get('interview_plus_rate_pct', 0.0)}%",
        f"- Fresh jobs (posted <=7 days): {summary.get('fresh_jobs_7d_pct', 0.0)}%",
        f"- Aged jobs (posted >14 days): {summary.get('aged_jobs_14d_pct', 0.0)}%",
        f"- Posted-date coverage: {summary.get('posted_date_coverage_pct', 0.0)}%",
        f"- Avg days since posted: {summary.get('avg_days_since_posted', 0.0)}",
        "",
        "## Status Breakdown",
        "",
    ]
    for status, count in summary["counts"].items():
        lines.append(f"- {status}: {count}")

    lines.extend(
        [
            "",
            "## Roles",
            "",
            "| ID | Match | Posted | Company | Title | Role | Missing | Apply URL |",
            "|---:|---:|---|---|---|---|---|---|",
        ]
    )
    for row in export_rows:
        lines.append(
            "| {pick_id} | {match} | {posted} | {company} | {title} | {role_lock} | {missing} | {url} |".format(
                pick_id=row.get("pick_id", ""),
                match=row.get("match_score", ""),
                posted=normalize_posted_date(row.get("job_posted_date", "")) or "-",
                company=str(row.get("company", "")).replace("|", "/"),
                title=str(row.get("title", "")).replace("|", "/"),
                role_lock=row.get("role_lock", "") or "-",
                missing=_compact_keywords(row.get("missing_tier1_keywords") or row.get("missing_keywords", "")),
                url=row.get("application_url", "") or row.get("job_url", ""),
            )
        )
    return "\n".join(lines) + "\n"


def sync_tracker_outputs(rows: List[dict]) -> List[dict]:
    final_rows = dedupe_and_normalize_rows(rows)
    write_xlsx(final_rows, AUDIT_XLSX_FILE)
    write_csv(final_rows, TRACKER_CSV_FILE)
    TRACKER_MD_FILE.parent.mkdir(parents=True, exist_ok=True)
    TRACKER_MD_FILE.write_text(build_tracker_markdown(final_rows), encoding="utf-8")
    return final_rows


def _value_has_signal(value: object) -> bool:
    text = str(value or "").strip()
    return text not in {"", "None", "nan", "null", "[]", "{}"}


def prepare_export_rows(rows: List[dict]) -> List[dict]:
    prepared_rows: List[dict] = []
    for idx, row in enumerate(rows, start=1):
        prepared = dict(row)
        prepared["pick_id"] = prepared.get("pick_id") or str(1000 + idx)
        prepared["posted_priority"] = posted_priority_bucket(prepared.get("job_posted_date", ""))
        prepared["days_since_posted"] = days_since_posted(prepared.get("job_posted_date", ""))
        prepared["application_url"] = prepared.get("application_url") or prepared.get("job_url", "")
        prepared_rows.append(prepared)
    return prepared_rows


def resolve_export_headers(rows: List[dict], preferred_headers: List[str], always_keep: Optional[Set[str]] = None) -> List[str]:
    keep = always_keep or set()
    headers: List[str] = []
    for header in preferred_headers:
        if header in keep or any(_value_has_signal(row.get(header, "")) for row in rows):
            headers.append(header)
    return headers


def build_extra_headers(rows: List[dict], base_headers: List[str]) -> List[str]:
    seen = set(base_headers)
    extras: List[str] = []
    for row in rows:
        for key in row.keys():
            if key in seen:
                continue
            seen.add(key)
            if any(_value_has_signal(other.get(key, "")) for other in rows):
                extras.append(key)
    return extras


def _compact_keywords(raw: object, limit: int = 3) -> str:
    terms = [str(x).strip() for x in str(raw or "").split(",") if str(x).strip()]
    filtered = [term for term in terms if not keyword_is_nonessential(term)]
    if not filtered:
        return "-"
    return ", ".join(filtered[:limit])


def read_settings() -> dict:
    with SETTINGS_FILE.open("r", encoding="utf-8") as f:
        return yaml.safe_load(f)


def load_seek_search_strategy() -> dict:
    if not SEEK_STRATEGY_FILE.exists():
        raise FileNotFoundError(f"Missing Seek strategy config: {SEEK_STRATEGY_FILE}")
    with SEEK_STRATEGY_FILE.open("r", encoding="utf-8") as f:
        return yaml.safe_load(f) or {}


def normalize_fact_text(text: str) -> str:
    s = str(text or "").strip()
    if not s:
        return ""
    s = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", s)
    s = s.replace("**", "").replace("`", "")
    s = s.replace("|", " ").replace("->", " to ")
    s = s.replace("鈫", " to ").replace("馃", " ")
    s = s.encode("ascii", "ignore").decode("ascii")
    s = re.sub(r"\s+", " ", s).strip()
    return s[:260]


def append_fact_unique(bucket: List[str], text: str) -> None:
    fact = normalize_fact_text(text)
    if not fact:
        return
    key = re.sub(r"[^a-z0-9]+", " ", fact.lower()).strip()
    if not key:
        return
    existing = {re.sub(r"[^a-z0-9]+", " ", x.lower()).strip() for x in bucket}
    if key not in existing:
        bucket.append(fact)


def load_external_profile_facts() -> dict:
    if not EXTERNAL_FACTS_FILE.exists():
        return {"facts": []}
    try:
        with EXTERNAL_FACTS_FILE.open("r", encoding="utf-8") as f:
            data = json.load(f)
        if isinstance(data, dict):
            facts = data.get("facts", [])
            if isinstance(facts, list):
                data["facts"] = [normalize_fact_text(x) for x in facts if normalize_fact_text(x)]
                return data
    except Exception:
        pass
    return {"facts": []}


def collect_facts_from_digital_twin_json(path: Path) -> List[str]:
    facts: List[str] = []
    if not path.exists():
        return facts
    try:
        with path.open("r", encoding="utf-8") as f:
            data = json.load(f)
    except Exception:
        return facts

    def append_metrics(prefix: str, metrics: dict, limit: int = 4) -> None:
        if not isinstance(metrics, dict):
            return
        metric_pairs = []
        for k, v in list(metrics.items())[:limit]:
            key = str(k).replace("_", " ").strip()
            val = str(v).strip()
            if key and val:
                metric_pairs.append(f"{key}: {val}")
        if metric_pairs:
            append_fact_unique(facts, f"{prefix}: {', '.join(metric_pairs)}")

    personal = data.get("personal", {}) if isinstance(data, dict) else {}
    if isinstance(personal, dict):
        append_fact_unique(facts, personal.get("summary", ""))
        append_fact_unique(facts, personal.get("elevator_pitch", ""))
        contact = personal.get("contact", {})
        if isinstance(contact, dict):
            append_fact_unique(facts, f"Portfolio: {contact.get('portfolio', '')}")
            append_fact_unique(facts, f"LinkedIn: {contact.get('linkedin', '')}")
    salary = data.get("salary_location", {}) if isinstance(data, dict) else {}
    if isinstance(salary, dict):
        append_fact_unique(facts, f"Salary expectations: {salary.get('salary_expectations', '')}")
        append_fact_unique(facts, f"Work authorization: {salary.get('work_authorization', '')}")
        countries = salary.get("countries_worked", [])
        if isinstance(countries, list) and countries:
            append_fact_unique(facts, f"Countries worked: {', '.join(str(x) for x in countries[:6])}")
        locations = salary.get("location_preferences", [])
        if isinstance(locations, list) and locations:
            append_fact_unique(facts, f"Location preferences: {', '.join(str(x) for x in locations[:6])}")
        append_fact_unique(facts, f"Remote experience: {salary.get('remote_experience', '')}")
        append_fact_unique(facts, f"Travel availability: {salary.get('travel_availability', '')}")

    for exp in (data.get("experience", []) if isinstance(data, dict) else [])[:6]:
        if not isinstance(exp, dict):
            continue
        company = str(exp.get("company", "")).strip()
        title = str(exp.get("title", "")).strip()
        duration = str(exp.get("duration", "")).strip()
        append_fact_unique(facts, f"Experience: {title} at {company} ({duration})")
        append_fact_unique(facts, exp.get("company_context", ""))
        append_fact_unique(facts, exp.get("team_structure", ""))
        for item in (exp.get("technical_skills_used", []) or [])[:6]:
            append_fact_unique(facts, f"{company} tools: {item}")
        for item in (exp.get("leadership_examples", []) or [])[:3]:
            append_fact_unique(facts, item)
        for item in (exp.get("projects_delivered", []) or [])[:3]:
            append_fact_unique(facts, item)
        for ach in (exp.get("achievements_star", []) or [])[:4]:
            if not isinstance(ach, dict):
                continue
            category = str(ach.get("category", "")).strip()
            action = str(ach.get("action", "")).strip()
            if category and action:
                append_fact_unique(facts, f"{category}: {action}")
            append_fact_unique(facts, ach.get("result", ""))
            append_metrics(f"{company} metrics", ach.get("metrics", {}))

    education = data.get("education", {}) if isinstance(data, dict) else {}
    if isinstance(education, dict):
        current = education.get("current", {})
        if isinstance(current, dict):
            append_fact_unique(
                facts,
                f"Education: {current.get('degree', '')} at {current.get('university', '')}, expected graduation {current.get('expected_graduation', '')}",
            )
            append_fact_unique(facts, f"Academic focus: {current.get('academic_focus', '')}")
            for item in (current.get("relevant_coursework", []) or [])[:5]:
                append_fact_unique(facts, item)
            for item in (current.get("key_projects", []) or [])[:4]:
                append_fact_unique(facts, item)
        undergrad = education.get("undergraduate", {})
        if isinstance(undergrad, dict):
            append_fact_unique(
                facts,
                f"Education: {undergrad.get('degree', '')} at {undergrad.get('university', '')}",
            )

    skills = data.get("skills", {}) if isinstance(data, dict) else {}
    if isinstance(skills, dict):
        technical = skills.get("technical", {})
        if isinstance(technical, dict):
            for item in (technical.get("data_tools", []) or [])[:5]:
                append_fact_unique(facts, item)
            for item in (technical.get("financial_accounting", []) or [])[:5]:
                append_fact_unique(facts, item)
            for item in (technical.get("ai_ml", []) or [])[:4]:
                append_fact_unique(facts, item)
            for item in (technical.get("development_tools", []) or [])[:4]:
                append_fact_unique(facts, item)
            for skill_item in (technical.get("programming_languages", []) or [])[:4]:
                if not isinstance(skill_item, dict):
                    continue
                language = str(skill_item.get("language", "")).strip()
                proficiency = str(skill_item.get("proficiency", "")).strip()
                years = str(skill_item.get("years_experience", "")).strip()
                if language:
                    append_fact_unique(facts, f"{language}: {proficiency} proficiency, {years} years experience")
                for use_case in (skill_item.get("use_cases", []) or [])[:2]:
                    append_fact_unique(facts, f"{language} use case: {use_case}")
        for item in (skills.get("business_skills", []) or [])[:6]:
            append_fact_unique(facts, item)
        for cert in (skills.get("certifications", []) or [])[:4]:
            if not isinstance(cert, dict):
                continue
            append_fact_unique(
                facts,
                f"Certification: {cert.get('name', '')} from {cert.get('issuer', '')} ({cert.get('date', cert.get('expected_completion', ''))})",
            )
            for item in (cert.get("skills", []) or [])[:4]:
                append_fact_unique(facts, item)
        for award in (skills.get("awards_recognition", []) or [])[:4]:
            if not isinstance(award, dict):
                continue
            append_fact_unique(
                facts,
                f"Award: {award.get('title', '')} - {award.get('description', '')}",
            )

    for item in (data.get("unique_value_propositions", []) if isinstance(data, dict) else [])[:5]:
        append_fact_unique(facts, item)

    talking_points = data.get("key_talking_points", {}) if isinstance(data, dict) else {}
    if isinstance(talking_points, dict):
        for bucket in ["for_data_analyst_roles", "for_ai_ml_roles", "for_graduate_programs", "for_fintech_roles"]:
            for item in (talking_points.get(bucket, []) or [])[:4]:
                append_fact_unique(facts, item)

    return facts


def collect_facts_from_markdown(path: Path) -> List[str]:
    facts: List[str] = []
    if not path.exists():
        return facts
    try:
        lines = path.read_text(encoding="utf-8").splitlines()
    except Exception:
        return facts
    for line in lines:
        clean = normalize_fact_text(line)
        if not clean:
            continue
        if clean.startswith("#"):
            continue
        if len(clean) < 18:
            continue
        if re.search(r"(\d+%|\$\d+|\d+\+|gpa|master|bachelor|visa|accuracy|latency|roi|forecast)", clean, re.I):
            append_fact_unique(facts, re.sub(r"^[\-\*\| ]+", "", clean))
    return facts


def build_external_profile_facts(source_dir: Path, max_facts: int = 80) -> dict:
    facts: List[str] = []
    json_path = source_dir / "digitaltwin.json"
    profile_md = source_dir / "DOUGLAS_MO_PROFESSIONAL_PROFILE.md"
    resume_md = source_dir / "archived" / "resume-materials" / "RESUME_DOUGLAS_MO.md"
    linkedin_md = source_dir / "archived" / "resume-materials" / "LINKEDIN_PROFILE_CONTENT.md"

    for item in collect_facts_from_digital_twin_json(json_path):
        append_fact_unique(facts, item)
    for md_path in [profile_md, resume_md, linkedin_md]:
        for item in collect_facts_from_markdown(md_path):
            append_fact_unique(facts, item)

    return {
        "source_dir": str(source_dir),
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "precedence_rule": "AI job config has priority when fields conflict; external facts are supplemental evidence only.",
        "facts": facts[:max_facts],
    }


def hydrate_profile_from_external(source_dir: str, max_facts: int = 80) -> None:
    src = Path(source_dir)
    if not src.exists():
        raise FileNotFoundError(f"Source directory not found: {src}")
    payload = build_external_profile_facts(src, max_facts=max_facts)
    EXTERNAL_FACTS_FILE.parent.mkdir(parents=True, exist_ok=True)
    with EXTERNAL_FACTS_FILE.open("w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False, indent=2)
    print(f"External facts updated: {EXTERNAL_FACTS_FILE}")
    print(f"Facts captured: {len(payload.get('facts', []))}")


def ensure_playwright_ready() -> None:
    if not PLAYWRIGHT_AVAILABLE:
        raise RuntimeError(
            "playwright is not installed. Install dependencies first: "
            "py -m pip install -r requirements.txt"
        )
    # Use a stable ASCII path for browser binaries to avoid Unicode-path lookup failures.
    os.environ.setdefault("PLAYWRIGHT_BROWSERS_PATH", str(PLAYWRIGHT_BROWSERS_DIR))
    PLAYWRIGHT_BROWSERS_DIR.mkdir(parents=True, exist_ok=True)


def resolve_playwright_chromium_executable(headless: bool = True) -> str:
    base = Path(os.environ.get("PLAYWRIGHT_BROWSERS_PATH", "")).expanduser()
    if not base.exists():
        return ""
    patterns = []
    if headless:
        patterns.extend(
            [
                "chromium_headless_shell-*/chrome-win/headless_shell.exe",
                "chromium-*/chrome-win/chrome.exe",
            ]
        )
    else:
        patterns.extend(
            [
                "chromium-*/chrome-win/chrome.exe",
                "chromium_headless_shell-*/chrome-win/headless_shell.exe",
            ]
        )
    for pat in patterns:
        matches = sorted(base.glob(pat), key=lambda p: p.stat().st_mtime, reverse=True)
        if matches:
            return str(matches[0])
    return ""


def launch_chromium_with_fallback(p, *, headless: bool, channel: Optional[str] = None):
    launch_timeout_ms = int(os.getenv("PLAYWRIGHT_LAUNCH_TIMEOUT_MS", "20000"))
    exe = resolve_playwright_chromium_executable(headless=headless)
    if exe:
        return p.chromium.launch(executable_path=exe, headless=headless, timeout=launch_timeout_ms)
    try:
        if channel:
            return p.chromium.launch(channel=channel, headless=headless, timeout=launch_timeout_ms)
        return p.chromium.launch(headless=headless, timeout=launch_timeout_ms)
    except Exception as exc:
        message = str(exc).lower()
        if "executable doesn't exist" not in message:
            raise
        exe = resolve_playwright_chromium_executable(headless=headless)
        if not exe:
            exe = resolve_playwright_chromium_executable(headless=not headless)
        if not exe:
            raise
        return p.chromium.launch(executable_path=exe, headless=headless, timeout=launch_timeout_ms)


def load_candidate_profile() -> dict:
    default_profile = {
        "name": "Douglas Mo",
        "email": "d157156@gmail.com",
        "phone": "0434 001 262",
        "linkedin": "https://www.linkedin.com/in/douglas-mo-67344531b/",
        "github": "https://github.com/DouglasMooooo",
        "portfolio": "https://douglasmo.vercel.app/",
        "gender": "Male",
        "ethnicity": "Chinese",
        "languages": ["English", "Mandarin", "Cantonese"],
        "address_line": "Algester Rd, Brisbane, QLD",
        "has_disability": "No",
        "is_lgbti": "No",
        "availability_note": "Available to start immediately for part-time or internship roles. Full-time from July 2026.",
    }
    if PROFILE_FILE.exists():
        try:
            with PROFILE_FILE.open("r", encoding="utf-8") as f:
                user_profile = json.load(f)
            if isinstance(user_profile, dict):
                default_profile.update(user_profile)
        except Exception:
            pass
    external_facts = load_external_profile_facts().get("facts", [])
    if isinstance(external_facts, list) and external_facts:
        default_profile["external_verified_facts"] = external_facts
    return default_profile


def get_verified_fact_snippets(profile: dict) -> List[str]:
    facts = []
    for k in ["name", "title", "location", "email", "phone", "work_authorization", "availability_note"]:
        v = str(profile.get(k, "")).strip()
        if v:
            facts.append(v)
    for item in profile.get("headline_achievements", []) or []:
        item = str(item).strip()
        if item:
            facts.append(item)
    for item in profile.get("finance_capabilities", []) or []:
        if isinstance(item, dict):
            evidence = str(item.get("evidence", "")).strip().lower()
            wording = str(item.get("wording", "")).strip()
            if wording and evidence in {"strong", "moderate"}:
                append_fact_unique(facts, wording)
    for item in profile.get("accounting_software_exposure", []) or []:
        if isinstance(item, dict):
            evidence = str(item.get("evidence", "")).strip().lower()
            wording = str(item.get("wording", "")).strip()
            if wording and evidence in {"strong", "moderate"}:
                append_fact_unique(facts, wording)
    for item in profile.get("external_verified_facts", []) or []:
        item = str(item).strip()
        if item:
            append_fact_unique(facts, item)
    return facts


def ensure_dirs() -> None:
    for path in [
        ROOT / "resumes",
        ROOT / "jobs",
        ROOT / "outputs",
        OUTPUT_APPLICATION_DIR,
        ROOT / "outputs" / "logs",
        PROFILE_DIR,
    ]:
        path.mkdir(parents=True, exist_ok=True)


def init_files() -> None:
    ensure_dirs()

    template_1 = ROOT / "resumes" / "resume_template_1.md"
    template_2 = ROOT / "resumes" / "resume_template_2.md"
    jobs_csv = JOBS_FILE
    portals_yaml = PORTALS_FILE

    if not template_1.exists():
        template_1.write_text(
            textwrap.dedent(
                """\
                # {name}
                Email: {email} | Phone: {phone} | LinkedIn: {linkedin}

                ## Professional Summary
                [Replace with your summary]

                ## Core Skills
                - [Skill A]
                - [Skill B]
                - [Skill C]

                ## Experience
                ### [Company] - [Role]
                - [Impact bullet with numbers]
                - [Impact bullet with numbers]

                ## Education
                - [Degree], [School]

                ## Certifications
                - [Certification]
                """
            ),
            encoding="utf-8",
        )

    if not template_2.exists():
        template_2.write_text(
            textwrap.dedent(
                """\
                # {name}
                {phone} | {email} | {city}

                ## Profile
                [Short profile for alternative role track]

                ## Technical Stack
                - [Tool/Language]
                - [Tool/Language]

                ## Selected Projects
                ### [Project Name]
                - [What was built]
                - [Business value]

                ## Work History
                ### [Company] - [Role]
                - [Result with KPI]

                ## Education
                - [Degree], [University]
                """
            ),
            encoding="utf-8",
        )

    if not jobs_csv.exists():
        jobs_csv.write_text(
            "job_id,platform,company,title,location,job_url,application_url,job_posted_date,jd_text,notes\n",
            encoding="utf-8",
        )
    if not portals_yaml.exists():
        portals_yaml.write_text(
            textwrap.dedent(
                """\
                title_filter:
                  positive:
                    - analyst
                    - data
                    - business
                    - ai
                    - graduate
                  negative:
                    - senior
                    - principal
                    - director
                    - manager
                tracked_companies:
                  - name: OpenAI
                    enabled: false
                    careers_url: https://boards.greenhouse.io/openai
                  - name: Anthropic
                    enabled: false
                    careers_url: https://jobs.ashbyhq.com/Anthropic
                  - name: Canva
                    enabled: false
                    careers_url: https://www.lifeatcanva.com/en/jobs
                """
            ),
            encoding="utf-8",
        )


def load_jobs() -> List[JobRecord]:
    if not JOBS_FILE.exists():
        raise FileNotFoundError(f"Missing jobs file: {JOBS_FILE}")

    rows: List[JobRecord] = []
    with JOBS_FILE.open("r", encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        for r in reader:
            rows.append(
                JobRecord(
                    job_id=(r.get("job_id") or "").strip(),
                    platform=(r.get("platform") or "").strip().lower(),
                    company=(r.get("company") or "").strip(),
                    title=(r.get("title") or "").strip(),
                    location=(r.get("location") or "").strip(),
                    job_url=(r.get("job_url") or "").strip(),
                    application_url=(r.get("application_url") or r.get("job_url") or "").strip(),
                    job_posted_date=normalize_posted_date((r.get("job_posted_date") or "").strip()),
                    jd_text=(r.get("jd_text") or "").strip(),
                    notes=(r.get("notes") or "").strip(),
                )
            )
    return rows


def load_templates(settings: dict) -> List[Tuple[str, str, Path]]:
    templates: List[Tuple[str, str, Path]] = []
    for rel in settings.get("template_files", []):
        p = ROOT / rel
        if p.exists():
            text = read_template_text(p)
            if text.strip():
                templates.append((p.name, text, p))
    if not templates:
        raise FileNotFoundError("No templates found under resumes/")
    return templates


def read_template_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in {".md", ".txt"}:
        return path.read_text(encoding="utf-8")
    if suffix == ".docx":
        doc = Document(str(path))
        lines = [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]
        return "\n".join(lines)
    return ""


def parse_base_resume_text(text: str) -> dict:
    data = {
        "name": "",
        "title": "",
        "contact": "",
        "links": "",
        "summary": [],
        "skills": [],
        "experience": [],
        "education": [],
        "projects": [],
        "availability": [],
    }
    lines = [ln.strip() for ln in str(text or "").splitlines() if ln.strip()]
    if not lines:
        return data

    if len(lines) > 0:
        data["name"] = lines[0]
    if len(lines) > 1:
        data["title"] = lines[1]
    if len(lines) > 2:
        data["contact"] = lines[2]
    if len(lines) > 3 and "linkedin" in lines[3].lower():
        data["links"] = lines[3]

    section = None
    current_role = None
    for line in lines[4:]:
        upper = line.upper()
        if upper in {"SUMMARY", "SKILLS", "EXPERIENCE", "EDUCATION", "PROJECTS", "AVAILABILITY"}:
            section = upper.lower()
            current_role = None
            continue

        if section == "summary":
            data["summary"].append(line)
        elif section == "skills":
            data["skills"].append(line)
        elif section == "experience":
            if "|" in line and not line.startswith("-"):
                current_role = {"header": line, "bullets": []}
                data["experience"].append(current_role)
            elif current_role is not None:
                bullet = line[2:].strip() if line.startswith("- ") else line
                current_role["bullets"].append(bullet)
        elif section == "education":
            data["education"].append(line)
        elif section == "projects":
            data["projects"].append(line)
        elif section == "availability":
            data["availability"].append(line)
    return data


def extract_text_from_url(url: str, timeout_s: float = 20.0) -> str:
    if not url:
        return ""
    headers = {
        "User-Agent": (
            "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
            "(KHTML, like Gecko) Chrome/125.0 Safari/537.36"
        )
    }
    try:
        with httpx.Client(timeout=timeout_s, follow_redirects=True, headers=headers) as client:
            resp = client.get(url)
            resp.raise_for_status()
        soup = BeautifulSoup(resp.text, "html.parser")
        for tag in soup(["script", "style", "noscript"]):
            tag.extract()
        text = re.sub(r"\s+", " ", soup.get_text(" ", strip=True))
        return text[:20000]
    except Exception:
        return ""


def login_platform(
    platform: str,
    headless: bool = False,
    wait_seconds: int = 120,
    channel: str = "msedge",
    use_system_edge: bool = False,
    cdp_port: int = 9222,
) -> None:
    ensure_playwright_ready()
    platform = platform.lower()
    if platform not in {"seek", "linkedin"}:
        raise ValueError("platform must be seek or linkedin")

    profile_path = PROFILE_DIR / platform
    profile_path.mkdir(parents=True, exist_ok=True)
    login_url = "https://www.seek.com.au/" if platform == "seek" else "https://www.linkedin.com/"

    if use_system_edge:
        ensure_system_edge_cdp(f"http://127.0.0.1:{cdp_port}", start_url=login_url)
        print(f"[{platform}] System Edge is ready on CDP port {cdp_port}.")
        print(f"Please confirm login in your normal Edge profile within {wait_seconds} seconds.")
        time.sleep(wait_seconds)
        exported = export_system_edge_session(platform, cdp_port=cdp_port, headless=headless, channel=channel)
        if exported:
            print(f"[{platform}] Session imported into stable local profile: {profile_path}")
        else:
            print(f"[{platform}] Could not import session from system Edge.")
        return

    with sync_playwright() as p:
        context = p.chromium.launch_persistent_context(
            user_data_dir=str(profile_path),
            headless=headless,
            viewport={"width": 1440, "height": 900},
            channel=channel,
            args=["--disable-blink-features=AutomationControlled"],
        )
        try:
            page = context.new_page()
            page.goto(login_url, wait_until="domcontentloaded")
            print(f"[{platform}] Browser opened. Please login manually.")
            print(f"Waiting {wait_seconds} seconds for manual login, then session will auto-save.")
            time.sleep(wait_seconds)
        except PlaywrightError:
            # User may close browser early; persistent profile may still contain login state.
            print(f"[{platform}] Browser closed before wait completed, attempting to save session.")
        try:
            context.storage_state(path=str(profile_path / "storage_state.json"))
        except PlaywrightError:
            pass
        try:
            context.close()
        except PlaywrightError:
            pass
        print(f"[{platform}] Session save attempt completed: {profile_path}")


def storage_state_path(platform: str) -> Path:
    return PROFILE_DIR / platform.lower() / "storage_state.json"


def export_system_edge_session(
    platform: str,
    *,
    cdp_port: int = 9222,
    headless: bool = False,
    channel: str = "msedge",
) -> bool:
    platform = platform.lower()
    profile_path = PROFILE_DIR / platform
    profile_path.mkdir(parents=True, exist_ok=True)
    state_path = storage_state_path(platform)
    cdp_url = f"http://127.0.0.1:{cdp_port}"

    with sync_playwright() as p:
        browser = None
        context = None
        seed_browser = None
        seed_context = None
        seed_page = None
        try:
            browser = p.chromium.connect_over_cdp(cdp_url)
            context = browser.contexts[0] if browser.contexts else browser.new_context()
            context.storage_state(path=str(state_path))
            cookies = context.cookies()

            seed_browser = launch_chromium_with_fallback(p, channel=channel, headless=headless)
            seed_context = seed_browser.new_context(storage_state=str(state_path), viewport={"width": 1440, "height": 900})
            if cookies:
                try:
                    seed_context.add_cookies(cookies)
                except Exception:
                    pass
            seed_page = seed_context.new_page()
            seed_page.goto("https://www.seek.com.au/" if platform == "seek" else "https://www.linkedin.com/", wait_until="domcontentloaded", timeout=45000)
            seed_context.storage_state(path=str(state_path))
            return state_path.exists()
        except Exception:
            return False
        finally:
            for obj in (seed_page, seed_context, seed_browser, context, browser):
                try:
                    if obj is not None:
                        obj.close()
                except Exception:
                    pass


def open_saved_platform_context(
    p,
    platform: str,
    *,
    headless: bool = False,
    channel: str = "msedge",
    prefer_headless_saved_state: bool = False,
) -> Tuple[object, object, str]:
    platform = platform.lower()
    profile_path = PROFILE_DIR / platform
    state_path = storage_state_path(platform)
    if state_path.exists():
        launch_headless = True if prefer_headless_saved_state else headless
        browser = launch_chromium_with_fallback(p, channel=channel, headless=launch_headless)
        context = browser.new_context(
            storage_state=str(state_path),
            viewport={"width": 1440, "height": 900},
            user_agent=(
                "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                "AppleWebKit/537.36 (KHTML, like Gecko) "
                "Chrome/125.0 Safari/537.36"
            ),
        )
        return browser, context, "storage_state"
    if profile_path.exists():
        context = p.chromium.launch_persistent_context(
            user_data_dir=str(profile_path),
            headless=headless,
            viewport={"width": 1440, "height": 900},
            channel=channel,
            args=["--disable-blink-features=AutomationControlled"],
        )
        return None, context, "persistent_profile"
    browser = launch_chromium_with_fallback(p, channel=channel, headless=headless)
    context = browser.new_context()
    return browser, context, "fresh"


def tokenize(text: str) -> List[str]:
    return re.findall(r"[A-Za-z][A-Za-z0-9\+\#\.\-]{1,}", text or "")


JD_NOISE_LINE_PATTERNS = [
    r"\bbenefits?\b",
    r"\bperks?\b",
    r"\bwhy\s+join\b",
    r"\babout\s+us\b",
    r"\bour\s+culture\b",
    r"\bwhat\s+we\s+offer\b",
    r"\bemployee\s+assistance\b",
    r"\bwellbeing\b",
    r"\bdiversity\b",
    r"\binclusion\b",
    r"\bequal\s+opportunity\b",
    r"\bhow\s+to\s+apply\b",
    r"\bsalary\b",
    r"\bremuneration\b",
    r"\bpackage\b",
]

JD_REQUIRED_PATTERNS = [
    r"\brequired\b",
    r"\bmust\s+have\b",
    r"\bessential\b",
    r"\byou\s+will\b",
    r"\bresponsibilit(y|ies)\b",
    r"\bexperience\s+with\b",
    r"\bproficien(t|cy)\b",
    r"\bqualification(s)?\b",
]

JD_PREFERRED_PATTERNS = [
    r"\bpreferred\b",
    r"\bnice\s+to\s+have\b",
    r"\bdesirable\b",
]

NON_ESSENTIAL_KEYWORDS = {
    "benefits", "benefit", "perks", "culture", "wellbeing", "well-being", "salary", "package",
    "bonus", "discounts", "leave", "insurance", "lunch", "snacks", "gym", "allowance",
}


def clean_jd_for_keywording(jd_text: str) -> str:
    lines = []
    for raw_line in str(jd_text or "").splitlines():
        line = re.sub(r"\s+", " ", raw_line).strip()
        if not line:
            continue
        low = line.lower()
        if any(re.search(p, low) for p in JD_NOISE_LINE_PATTERNS):
            continue
        lines.append(line)
    return "\n".join(lines)


def normalize_keyword_key(value: str) -> str:
    return re.sub(r"[^a-z0-9\+\#]+", " ", str(value or "").lower()).strip()


def keyword_is_nonessential(value: str) -> bool:
    k = normalize_keyword_key(value)
    if not k:
        return True
    if k in NON_ESSENTIAL_KEYWORDS:
        return True
    return any(x in k for x in ["benefit", "perk", "wellbeing", "culture fit", "salary"])


def candidate_location_keys(profile: dict) -> Set[str]:
    keys: Set[str] = set()
    for item in profile.get("location_preferences", []) or []:
        k = normalize_keyword_key(item)
        if k:
            keys.add(k)
    base = normalize_keyword_key(profile.get("location", ""))
    if base:
        keys.add(base)
    return keys


def keyword_is_location(value: str) -> bool:
    k = normalize_keyword_key(value)
    if not k:
        return False
    location_markers = ["brisbane", "melbourne", "sydney", "australia", "remote", "hybrid", "onsite", "on site"]
    return any(m in k for m in location_markers)


def infer_role_family(job_title: str, jd_text: str) -> str:
    text = f"{job_title} {jd_text}".lower()
    if any(k in text for k in ["accountant", "finance", "audit", "financial"]):
        return "finance"
    if any(re.search(p, text) for p in [r"\bdata analyst\b", r"\banalytics?\b", r"\bdashboard(s)?\b", r"\bsql\b", r"\bpython\b", r"\bpower bi\b"]):
        return "data_analyst"
    if any(re.search(p, text) for p in [r"\bsales\b", r"\baccount executive\b", r"\bbusiness development\b", r"\bbdm\b", r"\bclient(s)?\b"]):
        return "sales"
    if any(re.search(p, text) for p in [r"\bconsulting\b", r"\bconsultant\b", r"\badvisory\b"]):
        return "consulting"
    return "general"


def classify_keyword_tier(keyword: str, role_family: str) -> int:
    k = normalize_keyword_key(keyword)
    if not k or keyword_is_nonessential(k):
        return 4
    tier1_terms = {
        "sql", "python", "excel", "power bi", "tableau", "stakeholder management",
        "financial analysis", "financial reporting", "forecasting", "budgeting", "audit",
        "accounts payable", "accounts receivable", "reconciliation", "compliance", "erp",
        "sap", "xero", "myob", "client management", "negotiation", "revenue",
        "problem solving", "communication",
    }
    tier2_terms = {
        "reporting", "forecast", "analysis", "analyst", "process improvement", "modelling",
        "dashboarding", "data analysis", "cross functional", "documentation", "coordination",
        "variance analysis", "journal entries", "month end", "general ledger",
    }
    tier3_terms = {
        "manufacturing", "fintech", "banking", "supply chain", "retail", "esg", "saas",
        "commercial", "operations",
    }
    if k in tier1_terms:
        return 1
    if any(x in k for x in ["sql", "python", "power bi", "excel", "stakeholder", "audit", "forecast", "budget", "reconcil", "negotiat", "revenue", "client"]):
        return 1
    if k in tier2_terms:
        return 2
    if any(x in k for x in ["report", "analyse", "analysis", "process", "model", "dashboard", "journal", "month end", "ledger"]):
        return 2
    if k in tier3_terms:
        return 3
    if any(x in k for x in ["manufactur", "fintech", "supply chain", "banking", "retail", "commercial"]):
        return 3
    # Default to functional instead of noise when uncertain.
    return 2


def keyword_weight(tier: int, keyword: str, role_family: str) -> float:
    base = {1: 5.0, 2: 3.0, 3: 1.0, 4: 0.0}.get(tier, 0.0)
    k = normalize_keyword_key(keyword)
    if role_family == "data_analyst" and k in {"sql", "python", "power bi", "tableau"}:
        base *= 1.35
    if role_family == "sales" and any(x in k for x in ["client", "revenue", "negotiation"]):
        base *= 1.35
    if role_family == "consulting" and any(x in k for x in ["problem solving", "stakeholder"]):
        base *= 1.35
    return base


def has_meaningful_keyword_context(keyword: str, line: str) -> bool:
    l = normalize_keyword_key(line)
    k = normalize_keyword_key(keyword)
    if not k or k not in l:
        return False
    # Ignore bare keyword lists and comma stuffing.
    comma_count = line.count(",")
    if comma_count >= 5 and not re.search(r"\b(improved|reduced|built|led|delivered|managed|implemented|analysed|analyzed)\b", line.lower()):
        return False
    # Require a minimally descriptive line.
    word_count = len([w for w in l.split() if w])
    return word_count >= 5


def resume_context_hits(resume_text: str, keyword: str) -> Dict[str, bool]:
    lines = [ln.strip() for ln in resume_text.splitlines() if ln.strip()]
    section = ""
    in_experience = False
    hit_any = False
    hit_experience = False
    hit_measurable = False
    for line in lines:
        upper = line.upper().strip()
        if upper in {"SUMMARY", "SKILLS", "EXPERIENCE", "EDUCATION", "AVAILABILITY", "PROJECTS", "CERTIFICATIONS"}:
            section = upper
            in_experience = section == "EXPERIENCE"
            continue
        if not has_meaningful_keyword_context(keyword, line):
            continue
        hit_any = True
        if in_experience or line.startswith("-"):
            hit_experience = True
        if re.search(r"(\d+|%|\$|k\b|m\b|million|improved|reduced|increased|saved|growth)", line.lower()):
            hit_measurable = True
    return {
        "matched": hit_any,
        "experience": hit_experience,
        "measurable": hit_measurable,
    }


def keyword_stats(jd_text: str, resume_text: str, known_skills: List[str]) -> Dict[str, List[str]]:
    jd_tokens = {t.lower() for t in tokenize(jd_text)}
    resume_tokens = {t.lower() for t in tokenize(resume_text)}

    explicit = [k for k in known_skills if k.lower() in jd_tokens]
    covered = [k for k in explicit if k.lower() in resume_tokens]
    missing = [k for k in explicit if k.lower() not in resume_tokens]
    return {"covered": covered, "missing": missing}


def extract_top_jd_keywords(jd_text: str, limit: int = 24) -> List[str]:
    # Pull ATS-relevant keywords from requirement-heavy parts of JD and ignore noise sections.
    cleaned = clean_jd_for_keywording(jd_text)
    stop = {
        "with", "that", "this", "have", "will", "from", "your", "you", "our", "for",
        "and", "the", "are", "job", "role", "work", "team", "years", "year", "using",
        "experience", "skills", "ability", "required", "preferred", "about", "their",
        "into", "across", "through", "including", "within", "candidate",
        "http", "https", "www", "com", "au", "page", "visit", "apply", "currently",
        "working", "available", "company", "companies", "information", "more",
        "benefits", "perks", "culture", "salary", "package", "offer",
        "position", "opportunity", "successful", "applicant", "support", "supporting",
    }
    tokens = [t.lower() for t in tokenize(cleaned) if len(t) >= 3]
    # Phrase-first extraction for ATS signals common in AU entry-level roles.
    priority_phrases = [
        "financial reporting", "management reporting", "accounts payable", "accounts receivable",
        "bank reconciliation", "month end", "general ledger", "journal entries", "forecasting",
        "budgeting", "variance analysis", "audit", "tax", "bookkeeping", "compliance",
        "stakeholder management", "customer service", "data analysis", "power bi", "excel", "sql",
        "xero", "myob", "sap", "erp", "communication", "attention to detail", "problem solving",
    ]
    phrase_hits: List[str] = []
    cleaned_low = cleaned.lower()
    for phrase in priority_phrases:
        if phrase in cleaned_low and not keyword_is_nonessential(phrase):
            phrase_hits.append(phrase)

    freq: Dict[str, int] = {}
    for raw_line in cleaned.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        line_low = line.lower()
        weight = 1
        if any(re.search(p, line_low) for p in JD_REQUIRED_PATTERNS):
            weight = 3
        elif any(re.search(p, line_low) for p in JD_PREFERRED_PATTERNS):
            weight = 2
        line_tokens = [t.lower() for t in tokenize(line) if len(t) >= 3]
        for t in line_tokens:
            if re.search(r"\d", t):
                continue
            if "." in t or "/" in t:
                continue
            if t in stop:
                continue
            freq[t] = freq.get(t, 0) + weight
    for t in tokens:
        if re.search(r"\d", t):
            continue
        if "." in t or "/" in t:
            continue
        if t in stop:
            continue
        freq[t] = freq.get(t, 0) + 1
    ranked = sorted(freq.items(), key=lambda x: (-x[1], x[0]))
    merged: List[str] = []
    seen = set()
    for k in phrase_hits + [k for k, _ in ranked]:
        if keyword_is_nonessential(k):
            continue
        nk = normalize_keyword_key(k)
        if not nk or nk in seen:
            continue
        seen.add(nk)
        merged.append(k)
        if len(merged) >= limit:
            break
    return merged


def pick_template(templates: List[Tuple[str, str, Path]], jd_text: str) -> Tuple[str, str, Path]:
    # Always prefer the user's authored master resume when present.
    for name, text, path in templates:
        if path.name.lower() == "douglas_mo_resume_formatted.docx":
            return name, text, path
    jd_tokens = {t.lower() for t in tokenize(jd_text)}
    best_name = templates[0][0]
    best_text = templates[0][1]
    best_path = templates[0][2]
    best_score = -1
    for name, text, path in templates:
        tokens = {t.lower() for t in tokenize(text)}
        score = len(jd_tokens & tokens)
        if score > best_score:
            best_name, best_text, best_path, best_score = name, text, path, score
    return best_name, best_text, best_path


def detect_job_archetype(job: JobRecord, jd_text: str) -> str:
    text = f"{job.title} {jd_text}".lower()
    buckets = {
        "sales_customer": [
            "sales", "business development", "account manager", "customer service",
            "territory", "upsell", "client-facing", "retention",
        ],
        "finance_accounting": [
            "finance", "accounting", "financial", "bookkeeping", "tax", "audit", "fp&a",
        ],
        "data_analytics": [
            "data analyst", "analytics", "reporting", "bi", "power bi", "sql", "dashboard",
        ],
        "operations_admin": [
            "operations", "administration", "coordinator", "scheduling", "process", "workflow",
        ],
    }
    scores = {k: 0 for k in buckets}
    for k, keywords in buckets.items():
        for kw in keywords:
            if kw in text:
                scores[k] += 1
    best = max(scores.items(), key=lambda x: x[1])
    return best[0] if best[1] > 0 else "general"


def get_archetype_pack(archetype: str) -> dict:
    packs = {
        "sales_customer": {
            "headline": "Customer Growth & Commercial Execution",
            "summary_style": "customer-facing outcomes, pipeline growth, and relationship-led execution",
            "skills": [
                "Customer engagement, stakeholder communication, objection handling, solution selling",
                "Business development support, account growth, retention, cross-functional collaboration",
                "Excel, CRM workflows, KPI tracking, reporting, process discipline",
            ],
            "experience_intro": "Delivered customer and commercial outcomes through relationship management, structured follow-up, and data-informed decisions.",
        },
        "finance_accounting": {
            "headline": "Finance & Accounting Delivery",
            "summary_style": "financial control, reporting accuracy, and commercial insight",
            "skills": [
                "Financial reporting, reconciliation, compliance, audit readiness",
                "Budget tracking, forecast variance analysis, cost and profitability monitoring",
                "Excel (advanced), ERP exposure, Power BI, SQL, process controls",
            ],
            "experience_intro": "Led finance and reporting routines with strong accuracy standards and measurable business impact.",
        },
        "data_analytics": {
            "headline": "Data Analytics & Decision Support",
            "summary_style": "data analysis, insight generation, and measurable business decisions",
            "skills": [
                "Data analysis, KPI design, dashboarding, stakeholder reporting",
                "Root-cause analysis, forecasting, experimentation support, process improvement",
                "Power BI, Excel, SQL, Python, data storytelling",
            ],
            "experience_intro": "Built analytics outputs that improved prioritization, decision quality, and measurable business performance.",
        },
        "operations_admin": {
            "headline": "Operations & Administrative Reliability",
            "summary_style": "operational reliability, coordination quality, and service responsiveness",
            "skills": [
                "Administration, scheduling, documentation, process coordination",
                "Customer communication, issue resolution, cross-team execution",
                "Microsoft Office, Excel, reporting, workflow optimization",
            ],
            "experience_intro": "Executed operational workflows with high reliability, clear communication, and service discipline.",
        },
        "general": {
            "headline": "Business Operations & Analysis",
            "summary_style": "cross-functional execution, analytics support, and customer-centric delivery",
            "skills": [
                "Customer and stakeholder communication, collaboration, problem solving",
                "Business reporting, KPI tracking, process improvement, execution discipline",
                "Excel, Power BI, SQL, Python, Microsoft Office",
            ],
            "experience_intro": "Combined business operations and analytics skills to deliver measurable outcomes across teams.",
        },
    }
    return packs.get(archetype, packs["general"])


def build_role_locked_skill_categories(role_family: str, ats_keywords: List[str]) -> List[Tuple[str, List[str]]]:
    top_keywords = [k for k in ats_keywords if not keyword_is_nonessential(k)]
    kw_text = ", ".join(top_keywords[:6])
    if role_family == "finance":
        return [
            (
                "Accounting & Reporting",
                [
                    "Prepare and support financial reporting, general ledger work, reconciliations, month-end close, and tax compliance activities",
                    "Apply budgeting, forecasting, variance analysis, and financial controls to improve reporting quality and business visibility",
                ],
            ),
            (
                "Systems & Analysis",
                [
                    "Use Advanced Excel, Power BI, dashboard reporting, and data analysis to improve accuracy, efficiency, and decision support",
                    "Work with ERP and business systems including SAP, Microsoft Dynamics, Salesforce, and QuickBooks exposure",
                ],
            ),
            (
                "Client & Delivery",
                [
                    "Communicate clearly with clients and internal stakeholders while maintaining attention to detail and accuracy under deadlines",
                    f"Bring JD-aligned strengths in {kw_text}" if kw_text else "Bring JD-aligned strengths in tax, compliance, reporting, and communication",
                ],
            ),
        ]
    if role_family == "data_analyst":
        return [
            ("Data & Analytics", ["SQL, Python, Power BI, dashboard development, KPI reporting", "Data cleaning, exploratory analysis, forecasting, predictive modelling"]),
            ("Business Analysis", ["Stakeholder communication, requirements translation, problem solving, process improvement"]),
            ("Tools & Delivery", [f"JD-aligned strengths: {kw_text}" if kw_text else "JD-aligned strengths: SQL, Python, dashboards, reporting"]),
        ]
    return [
        ("Core Skills", [f"JD-aligned strengths: {kw_text}" if kw_text else "JD-aligned strengths: reporting, analysis, communication"]),
        ("Tools & Delivery", ["Excel, Power BI, SQL, stakeholder communication, process improvement"]),
    ]


def rewrite_bullet_for_role(role_header: str, bullet: str, role_family: str) -> str:
    text = str(bullet or "").strip()
    if not text:
        return text
    low_header = str(role_header or "").lower()
    low = text.lower()
    if role_family == "finance":
        replacements = {
            "Supported merchandising, replenishment, and checkout operations in a high-volume retail environment":
                "Supported high-volume retail operations with a focus on transaction accuracy, stock control, and reliable day-to-day execution",
            "Assisted stocktake and inventory verification, improving data accuracy and availability":
                "Assisted stocktake and inventory verification to improve record accuracy, stock visibility, and operational reliability",
            "Delivered customer service while maintaining transaction accuracy under time pressure":
                "Delivered customer service while maintaining transaction accuracy and attention to detail in a fast-paced environment",
            "Collaborated with cross-functional teams to ensure smooth store operations":
                "Worked with team members across store operations to resolve issues quickly and keep daily processes running smoothly",
            "Built AI-powered Digital Twin Interview Assistant achieving 95% accuracy and <2s latency":
                "Built workflow automation tools that improved information accuracy and reduced turnaround time in internal processes",
            "Designed RAG architecture using vector databases for scalable retrieval":
                "Designed structured retrieval workflows and system logic to improve information access and process consistency",
            "Developed AI agents and MCP-based tooling for workflow automation":
                "Developed automation tooling to support repeatable workflows, documentation quality, and operational efficiency",
            "Delivered 5+ AI solutions including ML models, RAG systems, and dashboards":
                "Delivered multiple internal solutions including dashboards and automation workflows that supported faster decision-making",
            "Optimised prompt and inference pipelines for performance":
                "Improved system performance and usability through testing, iteration, and process refinement",
            "Prepared financial statements under GAAP/IFRS standards":
                "Prepared financial statements under GAAP/IFRS standards to support accurate reporting and month-end review",
            "Managed general ledger, reconciliations, and month-end close":
                "Managed general ledger reconciliations and month-end close processes to maintain reporting accuracy and financial control",
            "Built dashboards identifying $1.8M at-risk revenue":
                "Built Excel and dashboard reporting that identified $1.8M in at-risk revenue and supported timely commercial action",
            "Reduced churn from 22% to 16%, protecting ~$600K revenue":
                "Analysed business performance trends and helped reduce churn from 22% to 16%, protecting approximately $600K in revenue",
            "Improved promotional ROI from 8% to 23% (+$85K)":
                "Evaluated promotional performance and improved ROI from 8% to 23%, generating an additional $85K in value",
            "Developed forecasting model improving accuracy from 70% to 94%":
                "Developed forecasting and reporting analysis that improved accuracy from 70% to 94% for better planning decisions",
            "Maintained 100% tax compliance":
                "Maintained 100% tax compliance through accurate filing support, documentation, and process discipline",
            "Processed 1,000+ corporate tax filings with full compliance":
                "Processed more than 1,000 corporate tax filings with full compliance and strong attention to detail",
            "Managed financial records for 20+ clients":
                "Managed financial records for 20+ clients, maintaining accuracy, organisation, timely follow-up, and clear client communication",
            "Reduced reporting time by 35% via automation":
                "Improved reporting efficiency by 35% through process automation and Excel-based workflow improvements",
            "Trained junior staff":
                "Supported and trained junior staff on process standards, documentation, and quality expectations",
        }
        if text in replacements:
            return replacements[text]
        if "account" in low_header or "tax" in low_header:
            return text
    return text


def build_resume_without_llm(
    template_text: str,
    job: JobRecord,
    covered: List[str],
    missing: List[str],
    profile: dict,
    ats_keywords: List[str],
    archetype: str,
    refinement_round: int = 0,
) -> str:
    base = parse_base_resume_text(template_text)
    name = base.get("name") or profile.get("name", "Douglas Mo")
    role_family = infer_role_family(job.title, job.jd_text)
    title = job.title if role_family == "finance" else (base.get("title") or profile.get("title", "Business Analyst"))
    contact = base.get("contact") or f"{profile.get('location','Brisbane, Australia')} | {profile.get('email','')} | {profile.get('phone','')}"
    links = "LinkedIn | GitHub | Portfolio"
    pack = get_archetype_pack(archetype)
    top_keywords = [k for k in ats_keywords if not keyword_is_nonessential(k)][:8]
    kw_phrase = ", ".join(top_keywords[:5])

    if role_family == "finance":
        summary_lines = [
            "Graduate accountant with a foundation in financial accounting, tax compliance, and reporting, backed by hands-on experience across financial statements, reconciliations, and month-end support.",
            f"My background includes identifying $1.8M at-risk revenue, improving forecast accuracy from 70% to 94%, and maintaining 100% tax compliance. I am now focused on building my career in taxation and accounting support, using Excel, reporting analysis, attention to detail, and clear client communication to contribute from day one{f', with strengths in {kw_phrase}' if kw_phrase else ''}.",
        ]
    else:
        summary_lines = [
            "I have a background in finance, accounting, and business analytics, with hands-on experience supporting reporting, operational decisions, and measurable commercial outcomes.",
            f"For {job.title} opportunities, I bring practical strengths in {pack['summary_style']}{f', including {kw_phrase}' if kw_phrase else ''}.",
        ]
    if refinement_round > 0 and role_family == "finance":
        summary_lines[1] = (
            f"My experience includes financial reporting, general ledger support, reconciliations, forecasting, tax-related work, and Excel-based analysis, with measurable impact across revenue protection, reporting accuracy, and compliance outcomes"
            f"{f', including {kw_phrase}' if kw_phrase else ''}. I am targeting graduate accounting work where I can continue building practical taxation and client service capability."
        )

    skill_categories = build_role_locked_skill_categories(role_family, ats_keywords)

    availability = base.get("availability") or []
    if not availability:
        availability = [str(profile.get("availability_note", "")).strip()]
    if profile.get("relocation_willing"):
        prefs = ", ".join([str(x).strip() for x in (profile.get("location_preferences", []) or []) if str(x).strip()])
        if prefs and not any("relocation" in x.lower() for x in availability):
            availability.append(f"Open to relocation across: {prefs}")

    out: List[str] = [name, title, contact, links, "", "SUMMARY"]
    out.extend(summary_lines)
    out.extend(["", "SKILLS"])
    for heading, bullets in skill_categories:
        out.append(heading)
        for bullet in bullets:
            out.append(f"- {bullet}")

    out.extend(["", "EXPERIENCE"])
    for role in base.get("experience", []):
        header = role.get("header", "").strip()
        if not header:
            continue
        out.append(header)
        for bullet in role.get("bullets", []):
            txt = rewrite_bullet_for_role(header, bullet, role_family).strip()
            if txt:
                out.append(f"- {txt}")

    if base.get("education"):
        out.extend(["", "EDUCATION"])
        out.extend(base["education"])

    if availability:
        out.extend(["", "AVAILABILITY"])
        out.extend([x for x in availability if str(x).strip()])

    return "\n".join(out).strip() + "\n"


def build_cover_letter_without_llm(job: JobRecord, covered: List[str], missing: List[str], profile: dict) -> str:
    matched = ", ".join(covered[:8]) if covered else "relevant skills from the job description"
    gap_line = ", ".join(missing[:6]) if missing else "No obvious keyword gaps detected."
    verified_impact = (profile.get("headline_achievements") or [])[:2]
    impact_line = " ".join(verified_impact) if verified_impact else ""
    return textwrap.dedent(
        f"""\
        {datetime.now().strftime("%Y-%m-%d")}

        Hiring Manager
        {job.company}

        Re: {job.title}

        Dear Hiring Manager,

        I am writing to apply for the {job.title} opportunity at {job.company}. My background aligns strongly with the requirements in your job description, especially around {matched}.

        I focus on delivering measurable outcomes, clear communication, and reliable execution in cross-functional teams. I am confident I can add value quickly and contribute to your team in {job.location or "this role"}.

        Based on the role requirements, I have tailored my resume to emphasize the most relevant strengths while keeping claims accurate and evidence-based. {impact_line}

        {profile.get("availability_note", "")}

        Thank you for your time and consideration.

        Sincerely,
        Douglas Mo

        ---
        ATS alignment notes (internal):
        Missing keywords to consider adding in future updates: {gap_line}
        """
    )


def call_openai(prompt: str, model: str, api_key: str) -> str:
    # Uses Responses API directly to avoid SDK/version mismatch in local environments.
    url = "https://api.openai.com/v1/responses"
    payload = {
        "model": model,
        "input": prompt,
        "temperature": 0.2,
    }
    headers = {
        "Authorization": f"Bearer {api_key}",
        "Content-Type": "application/json",
    }
    with httpx.Client(timeout=120) as client:
        r = client.post(url, headers=headers, content=json.dumps(payload))
        r.raise_for_status()
        data = r.json()
    # Compatible fallback extraction:
    output_text = data.get("output_text")
    if isinstance(output_text, str) and output_text.strip():
        return output_text.strip()
    for item in data.get("output", []):
        for chunk in item.get("content", []):
            if chunk.get("type") == "output_text":
                txt = chunk.get("text", "")
                if txt.strip():
                    return txt.strip()
    return ""


def parse_resume_bundle(raw_text: str) -> dict:
    text = (raw_text or "").strip()
    if not text:
        return {}
    # Try direct JSON
    try:
        data = json.loads(text)
        if isinstance(data, dict):
            return data
    except Exception:
        pass
    # Try fenced JSON
    m = re.search(r"```json\s*(\{.*?\})\s*```", text, flags=re.DOTALL | re.IGNORECASE)
    if m:
        try:
            data = json.loads(m.group(1))
            if isinstance(data, dict):
                return data
        except Exception:
            pass
    # Try first object-like block
    m2 = re.search(r"(\{[\s\S]*\})", text)
    if m2:
        try:
            data = json.loads(m2.group(1))
            if isinstance(data, dict):
                return data
        except Exception:
            pass
    return {}


def build_prompt(
    template_text: str,
    jd_text: str,
    job: JobRecord,
    ats_keywords: List[str],
    verified_facts: List[str],
    archetype: str,
) -> str:
    preserved_roles = extract_role_headers_from_resume(template_text)
    return textwrap.dedent(
        f"""\
        You are a senior resume optimizer.
        Improve the existing resume, not rewrite it from scratch.

        CRITICAL RULES:
        - Do NOT remove or merge job roles.
        - Keep original job titles, company names, locations, and dates.
        - Never replace real roles with generic labels like "Relevant Experience".
        - Preserve measurable achievements and business outcomes.
        - Write in a natural professional tone; do NOT use robotic phrasing.
        - Integrate keywords in context (sentences/bullets), not keyword stuffing.
        - Remove low-value noise from emphasis: benefits, perks, employer branding, culture slogans.
        - Use standard sections only: SUMMARY, SKILLS, EXPERIENCE, EDUCATION, AVAILABILITY.
        - No tables, no graphics, no icons.
        - No fabricated experience, companies, dates, or metrics.
        - Do not delete content. If you simplify, rewrite rather than remove.
        - Preserve all original roles, companies, dates, and key achievements from the base resume.

        EXPERIENCE FORMAT (STRICT):
        - Use: Job Title | Company | Location | Dates
        - Then bullets with Action + What you did + Business outcome
        - Keep quantified impact wherever possible

        SKILLS FORMAT:
        - 3 to 4 clear categories
        - No "Role Keywords" section

        PRESERVED ROLE HEADERS FROM BASE RESUME (must keep all):
        {chr(10).join(f"- {x}" for x in preserved_roles) if preserved_roles else "- (none parsed; preserve all concrete roles from base resume text)"}

        Return STRICT JSON only with this schema:
        {{
          "role_lock": "string",
          "keywords_injected": ["..."],
          "ats_match_score": 0,
          "missing_keywords": ["..."],
          "risk_flags": ["..."],
          "improvements": ["..."],
          "remove_for_clarity": ["..."],
          "safe_version": "full resume text",
          "aggressive_version": "full resume text",
          "human_version": "full resume text"
        }}

        Job meta:
        - Title: {job.title}
        - Company: {job.company}
        - Location: {job.location}
        - Role archetype target: {archetype}
        - ATS target keywords: {", ".join(ats_keywords[:30])}

        JD:
        {jd_text[:12000]}

        Verified candidate facts:
        {chr(10).join(f"- {x}" for x in verified_facts[:50])}

        Base resume:
        {template_text}
        """
    )


def build_cover_letter_prompt(
    job: JobRecord,
    jd_text: str,
    resume_text: str,
    ats_keywords: List[str],
    verified_facts: List[str],
    archetype: str,
) -> str:
    return textwrap.dedent(
        f"""\
        You are a professional career writer.
        Write a tailored cover letter in markdown for this exact job.

        Constraints:
        - Keep it concise (220-320 words).
        - Tone: confident, specific, professional.
        - Use ATS-relevant wording naturally from JD.
        - No invented facts, achievements, or technologies.
        - Only reference achievements that appear in verified facts or resume text.
        - Use strong but truthful positioning, emphasizing transferable strengths and immediate business value.
        - Use this structure:
          1) opening with role + motivation
          2) 1-2 paragraphs on fit and impact
          3) closing + call to action
        - Do not include placeholders.

        Job meta:
        - Title: {job.title}
        - Company: {job.company}
        - Location: {job.location}
        - Role archetype target: {archetype}
        - ATS target keywords: {", ".join(ats_keywords[:20])}

        JD:
        {jd_text[:12000]}

        Resume:
        {resume_text[:12000]}

        Verified candidate facts:
        {chr(10).join(f"- {x}" for x in verified_facts[:40])}
        """
    )


def evaluate_resume(
    jd_text: str,
    resume_text: str,
    known_skills: List[str],
    candidate_profile: Optional[dict] = None,
    job_title: str = "",
) -> Dict[str, str]:
    profile = candidate_profile or {}
    lower_resume = resume_text.lower()
    role_family = infer_role_family(job_title, jd_text)
    ats_keywords = [k for k in extract_top_jd_keywords(jd_text, limit=28) if not keyword_is_nonessential(k)]
    location_keys = candidate_location_keys(profile)
    relocation_willing = bool(profile.get("relocation_willing"))

    matched_all: List[str] = []
    missing_all: List[str] = []
    matched_tier1: List[str] = []
    missing_tier1: List[str] = []
    tiered: List[Tuple[str, int, float]] = []
    for k in ats_keywords:
        tier = classify_keyword_tier(k, role_family)
        wt = keyword_weight(tier, k, role_family)
        if tier == 4 or wt <= 0:
            continue
        tiered.append((k, tier, wt))

    total_possible = 0.0
    matched_points = 0.0
    for k, tier, wt in tiered:
        total_possible += wt * 1.6  # includes possible context bonus ceiling
        key_norm = normalize_keyword_key(k)
        ctx = resume_context_hits(resume_text, k)
        in_resume = ctx["matched"]
        # If candidate is relocation-flexible, location keywords from preferred cities should not count as missing.
        if not in_resume and relocation_willing and keyword_is_location(k):
            if any((lk in key_norm) or (key_norm in lk) for lk in location_keys):
                in_resume = True
                ctx = {"matched": True, "experience": False, "measurable": False}
        if in_resume:
            bonus = 0.0
            if ctx.get("experience"):
                bonus += 0.3
            if ctx.get("measurable"):
                bonus += 0.3
            matched_points += wt * (1.0 + bonus)
            matched_all.append(k)
            if tier == 1:
                matched_tier1.append(k)
        else:
            missing_all.append(k)
            if tier == 1:
                missing_tier1.append(k)

    # 1) Weighted ATS keyword relevance score (70%)
    kw_score = round(100 * matched_points / max(total_possible, 1.0), 1)

    # 2) Standard sections (15%)
    section_targets = ["summary", "skills", "experience", "education", "availability"]
    present = sum(1 for s in section_targets if s in lower_resume)
    section_score = round(100 * present / len(section_targets), 1)

    # 3) Impact density: bullets with measurable outcomes (10%)
    bullets = [ln.strip() for ln in resume_text.splitlines() if ln.strip().startswith("-")]
    measurable = 0
    for b in bullets:
        if re.search(r"(\d+|%|\$|million|kpi|roi|latency|accuracy|saved|reduced|improved)", b.lower()):
            measurable += 1
    impact_score = round(100 * measurable / max(len(bullets), 1), 1)

    # 4) Role-title alignment (5%)
    title_source = f"{job_title} {jd_text[:220]}".strip()
    jd_title_tokens = [t.lower() for t in tokenize(title_source) if len(t) >= 4][:10]
    role_hits = sum(1 for t in jd_title_tokens if t in lower_resume)
    role_score = round(100 * role_hits / max(len(jd_title_tokens), 1), 1)

    ats_score = round(
        kw_score * 0.70 + section_score * 0.15 + impact_score * 0.10 + role_score * 0.05,
        1,
    )
    critical_gaps = ", ".join(missing_tier1[:6]) if missing_tier1 else "None"
    improvements: List[str] = []
    if missing_tier1:
        improvements.append(f"Add concise experience bullets that show: {', '.join(missing_tier1[:3])}.")
    if len([m for m in matched_all if m in matched_tier1]) < max(2, len(matched_tier1)):
        improvements.append("Move core tools/skills from SKILLS into EXPERIENCE bullets with business outcomes.")
    improvements.append("Keep each role with clear title/company/location/dates and quantify at least one impact bullet per role.")
    improvements = improvements[:3]

    return {
        "match_score": f"{ats_score}",
        "ats_score": f"{ats_score}",
        "section_score": f"{section_score}",
        "ats_keywords_used": ", ".join(matched_all[:20]) if matched_all else "",
        "covered_keywords": ", ".join(matched_all[:20]) if matched_all else "",
        "missing_keywords": ", ".join(missing_all[:15]) if missing_all else "None",
        "matched_tier1_keywords": ", ".join(matched_tier1[:12]) if matched_tier1 else "",
        "missing_tier1_keywords": ", ".join(missing_tier1[:12]) if missing_tier1 else "None",
        "critical_gaps": critical_gaps,
        "suggested_improvements": " | ".join(improvements),
    }


def extract_role_headers_from_resume(text: str) -> List[str]:
    headers: List[str] = []
    seen = set()
    for raw in str(text or "").splitlines():
        line = re.sub(r"\s+", " ", raw).strip()
        if not line:
            continue
        if "|" not in line:
            continue
        if not re.search(r"\b(19|20)\d{2}\b", line):
            continue
        key = normalize_keyword_key(line)
        if key and key not in seen:
            seen.add(key)
            headers.append(line)
    return headers


def contains_role_header(resume_text: str, header: str) -> bool:
    short = normalize_keyword_key("|".join([x.strip() for x in str(header).split("|")[:2]]))
    if not short:
        return True
    return short in normalize_keyword_key(resume_text)


def postprocess_resume_text(resume_text: str, profile: dict) -> str:
    lines = str(resume_text or "").splitlines()
    out: List[str] = []
    skip_next_keyword_line = False
    for ln in lines:
        clean = ln.strip()
        if not clean:
            out.append("")
            continue
        if "role keywords" in clean.lower():
            skip_next_keyword_line = True
            continue
        if skip_next_keyword_line and clean.startswith("-"):
            skip_next_keyword_line = False
            continue
        skip_next_keyword_line = False
        # Strip emoji/icons to keep ATS-safe plain text.
        clean = re.sub(r"[\U0001F300-\U0001FAFF]+", "", clean).strip()
        ln = clean
        if clean.lower().startswith("candidate aligned to"):
            rewritten = "I bring hands-on experience across business analysis, finance, and operational improvement, and I tailor my approach to the role requirements and team priorities."
            out.append(rewritten)
            continue
        out.append(ln)

    text = "\n".join(out).strip()
    if "AVAILABILITY" not in text.upper():
        availability = str(profile.get("availability_note", "")).strip()
        if availability:
            extra = ["", "AVAILABILITY", availability]
            if profile.get("relocation_willing"):
                prefs = ", ".join([str(x).strip() for x in (profile.get("location_preferences", []) or []) if str(x).strip()])
                if prefs:
                    extra.append(f"Open to relocation and hybrid opportunities across: {prefs}.")
            text = (text + "\n" + "\n".join(extra)).strip()
    return text + "\n"


def write_resume_docx_clean(text: str, path: Path, template_links: Dict[str, str], profile: Optional[dict] = None) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    style_doc_defaults(doc)

    section_headers = {
        "SUMMARY",
        "SKILLS",
        "EXPERIENCE",
        "EDUCATION",
        "AVAILABILITY",
    }

    lines = [ln.rstrip() for ln in str(text or "").splitlines()]
    idx = 0
    if lines:
        p = doc.add_paragraph(lines[0].strip())
        if p.runs:
            p.runs[0].bold = True
            p.runs[0].font.size = Pt(16)
        idx = 1
    if idx < len(lines):
        p = doc.add_paragraph(lines[idx].strip())
        if p.runs:
            p.runs[0].bold = True
            p.runs[0].font.size = Pt(11)
        idx += 1
    if idx < len(lines):
        doc.add_paragraph(lines[idx].strip())
        idx += 1

    link_map = template_links or {}
    if not link_map and profile:
        link_map = {
            "linkedin": profile.get("linkedin", "https://www.linkedin.com/"),
            "github": profile.get("github", "https://github.com/"),
            "portfolio": profile.get("portfolio", "https://"),
        }
    if link_map:
        p = doc.add_paragraph("")
        add_hyperlink(p, "LinkedIn", link_map.get("linkedin", "https://www.linkedin.com/"))
        p.add_run(" | ")
        add_hyperlink(p, "GitHub", link_map.get("github", "https://github.com/"))
        p.add_run(" | ")
        add_hyperlink(p, "Portfolio", link_map.get("portfolio", "https://"))

    for raw in lines[idx:]:
        clean = raw.strip()
        if not clean:
            doc.add_paragraph("")
            continue
        if "linkedin" in clean.lower() and "github" in clean.lower() and "portfolio" in clean.lower():
            continue
        if clean.upper() in section_headers:
            p = doc.add_paragraph(clean)
            if p.runs:
                p.runs[0].bold = True
                p.runs[0].font.size = Pt(11.5)
            continue
        if clean.startswith("- "):
            doc.add_paragraph(clean[2:].strip(), style="List Bullet")
            continue
        if "|" in clean and re.search(r"\b(19|20)\d{2}\b", clean):
            p = doc.add_paragraph(clean)
            if p.runs:
                p.runs[0].bold = True
            continue
        doc.add_paragraph(clean)
    return save_doc_with_fallback(doc, path)


def run_pipeline(use_openai: bool) -> None:
    ensure_dirs()
    load_dotenv(ROOT / ".env")
    settings = read_settings()
    templates = load_templates(settings)
    jobs = load_jobs()
    known_skills = settings.get("known_skills", [])
    candidate_profile = load_candidate_profile()
    verified_facts = get_verified_fact_snippets(candidate_profile)

    model = os.getenv("OPENAI_MODEL", settings.get("default_openai_model", "gpt-5-mini"))
    api_key = os.getenv("OPENAI_API_KEY", "")
    can_use_openai = use_openai and bool(api_key)

    existing_rows = read_xlsx_rows(AUDIT_XLSX_FILE)
    results = []
    for i, job in enumerate(jobs, start=1):
        job_id = job.job_id or f"job_{i:03d}"
        jd_text = job.jd_text.strip()
        if not jd_text and job.job_url:
            jd_text = extract_text_from_url(job.job_url)

        if not jd_text:
            results.append(
                {
                    "job_id": job_id,
                    "platform": job.platform,
                    "company": job.company,
                    "title": job.title,
                    "status": "skipped_no_jd",
                    "resume_file": "",
                    "resume_safe_file": "",
                    "resume_aggressive_file": "",
                    "resume_human_file": "",
                    "cover_letter_file": "",
                    "role_lock": "",
                    "keywords_injected": "",
                    "match_score": "",
                    "ats_match_score_llm": "",
                    "ats_score": "",
                    "section_score": "",
                    "ats_keywords_used": "",
                    "covered_keywords": "",
                    "missing_keywords": "",
                    "missing_keywords_llm": "",
                    "risk_flags": "",
                    "improvements": "",
                    "remove_for_clarity": "",
                    "job_url": job.job_url,
                    "application_url": job.application_url or job.job_url,
                    "job_posted_date": normalize_posted_date(job.job_posted_date),
                    "days_since_posted": days_since_posted(job.job_posted_date),
                    "timestamp": datetime.now().isoformat(timespec="seconds"),
                }
            )
            continue

        template_name, template_text, template_path = pick_template(templates, jd_text)
        template_links = extract_docx_field_links(template_path) if template_path.suffix.lower() == ".docx" else {}
        llm_caller = None
        if can_use_openai:
            llm_caller = lambda prompt: call_openai(prompt, model=model, api_key=api_key)

        generated = optimize_application(
            job_title=job.title,
            company=job.company,
            location=job.location,
            jd_text=jd_text,
            base_resume_text=template_text,
            candidate_profile=candidate_profile,
            verified_facts=verified_facts,
            llm_caller=llm_caller,
            score_threshold=85.0,
        )
        if generated.skip_reason:
            results.append(
                {
                    "job_id": job_id,
                    "platform": job.platform,
                    "company": job.company,
                    "title": job.title,
                    "status": "skipped",
                    "resume_file": "",
                    "resume_safe_file": "",
                    "resume_aggressive_file": "",
                    "resume_human_file": "",
                    "cover_letter_file": "",
                    "template_used": template_name,
                    "role_lock": generated.role_lock,
                    "keywords_injected": ", ".join(generated.jd_signals.keywords[:15]),
                    "match_score": "",
                    "ats_match_score_llm": "",
                    "ats_score": "",
                    "section_score": "",
                    "ats_keywords_used": "",
                    "covered_keywords": "",
                    "missing_keywords": "",
                    "matched_tier1_keywords": "",
                    "missing_tier1_keywords": "",
                    "critical_gaps": "",
                    "suggested_improvements": generated.skip_reason,
                    "missing_keywords_llm": "",
                    "risk_flags": "hard_eligibility_mismatch",
                    "improvements": generated.skip_reason,
                    "remove_for_clarity": "",
                    "strict_profile_facts_used": "",
                    "job_url": job.job_url,
                    "application_url": job.application_url or job.job_url,
                    "job_posted_date": normalize_posted_date(job.job_posted_date),
                    "days_since_posted": days_since_posted(job.job_posted_date),
                    "timestamp": datetime.now().isoformat(timespec="seconds"),
                }
            )
            continue
        human_resume_text = postprocess_resume_text(generated.resume_text, candidate_profile)
        cover_letter_text = generated.cover_letter_text
        evaluation = {
            "match_score": f"{generated.ats_result.score}",
            "ats_score": f"{generated.ats_result.score}",
            "section_score": f"{generated.ats_result.debug.get('section_score', 0.0)}",
            "ats_keywords_used": ", ".join(generated.ats_result.matched_keywords),
            "covered_keywords": ", ".join(generated.ats_result.matched_keywords),
            "missing_keywords": ", ".join(generated.ats_result.missing_keywords) if generated.ats_result.missing_keywords else "None",
            "matched_tier1_keywords": ", ".join(generated.ats_result.matched_tier1_keywords),
            "missing_tier1_keywords": ", ".join(generated.ats_result.missing_tier1_keywords) if generated.ats_result.missing_tier1_keywords else "None",
            "critical_gaps": ", ".join(generated.ats_result.missing_tier1_keywords) if generated.ats_result.missing_tier1_keywords else "None",
            "suggested_improvements": " | ".join(generated.ats_result.improvements),
        }
        out_file, cl_file = build_output_paths(job.company, job_id)
        out_file = write_resume_docx_clean(human_resume_text, out_file, template_links, candidate_profile)
        cl_file = write_cover_letter_docx(cover_letter_text, cl_file)

        results.append(
            {
                "job_id": job_id,
                "platform": job.platform,
                "company": job.company,
                "title": job.title,
                "status": "ready_to_apply",
                "resume_file": str(out_file),
                "resume_safe_file": "",
                "resume_aggressive_file": "",
                "resume_human_file": str(out_file),
                "cover_letter_file": str(cl_file),
                "template_used": template_name,
                "role_lock": generated.role_lock,
                "keywords_injected": ", ".join(generated.jd_signals.keywords[:15]),
                "match_score": evaluation["match_score"],
                "ats_match_score_llm": "",
                "ats_score": evaluation["ats_score"],
                "section_score": evaluation["section_score"],
                "ats_keywords_used": evaluation["ats_keywords_used"],
                "covered_keywords": evaluation["covered_keywords"],
                "missing_keywords": evaluation["missing_keywords"],
                "matched_tier1_keywords": evaluation["matched_tier1_keywords"],
                "missing_tier1_keywords": evaluation["missing_tier1_keywords"],
                "critical_gaps": evaluation["critical_gaps"],
                "suggested_improvements": evaluation["suggested_improvements"],
                "missing_keywords_llm": "",
                "risk_flags": "",
                "improvements": " | ".join(generated.ats_result.improvements),
                "remove_for_clarity": "",
                "strict_profile_facts_used": " | ".join(generated.selected_profile_facts[:8]),
                "job_url": job.job_url,
                "application_url": job.application_url or job.job_url,
                "job_posted_date": normalize_posted_date(job.job_posted_date),
                "days_since_posted": days_since_posted(job.job_posted_date),
                "timestamp": datetime.now().isoformat(timespec="seconds"),
            }
        )

    final_rows = sync_tracker_outputs(existing_rows + results)
    summary = summarize_tracker(final_rows)
    print(f"Updated audit file: {AUDIT_XLSX_FILE}")
    print(f"Tracker rows: {summary['total']} | average match score: {summary['avg_match_score']}")
    print(f"Tracker markdown: {TRACKER_MD_FILE}")


def safe_name(s: str) -> str:
    cleaned = re.sub(r"[^A-Za-z0-9\-_]+", "_", (s or "role")).strip("_")
    return cleaned[:60] if cleaned else "role"


def build_output_paths(company: str, job_id: str) -> Tuple[Path, Path]:
    base = safe_name(company) or "company"
    resume = OUTPUT_APPLICATION_DIR / f"{base}_resume.docx"
    cover = OUTPUT_APPLICATION_DIR / f"{base}_coverletter.docx"
    if resume.exists() or cover.exists():
        resume = OUTPUT_APPLICATION_DIR / f"{base}_{safe_name(job_id)}_resume.docx"
        cover = OUTPUT_APPLICATION_DIR / f"{base}_{safe_name(job_id)}_coverletter.docx"
    return resume, cover


def add_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1155CC")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(color)
    r_pr.append(underline)
    run.append(r_pr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def extract_docx_field_links(path: Path) -> Dict[str, str]:
    links: Dict[str, str] = {}
    try:
        with zipfile.ZipFile(path, "r") as zf:
            xml = zf.read("word/document.xml").decode("utf-8", errors="ignore")
    except Exception:
        return links

    field_pattern = re.compile(
        r'HYPERLINK\s+&quot;([^&]+)&quot;(.+?)w:fldChar w:fldCharType="end"',
        flags=re.DOTALL,
    )
    for m in field_pattern.finditer(xml):
        url = html.unescape(m.group(1)).strip()
        seg = m.group(2)
        texts = re.findall(r"<w:t[^>]*>(.*?)</w:t>", seg, flags=re.DOTALL)
        label = html.unescape("".join(texts)).strip()
        if not label or not url:
            continue
        links[label.lower()] = url
    return links


def style_doc_defaults(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)


def write_resume_docx(text: str, path: Path, template_links: Dict[str, str], profile: Optional[dict] = None) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    style_doc_defaults(doc)

    section_headers = {
        "SUMMARY",
        "SKILLS",
        "EXPERIENCE",
        "EDUCATION",
        "PROJECTS",
        "CERTIFICATIONS",
        "AVAILABILITY",
    }

    raw_lines = text.splitlines()
    lines: List[str] = []
    for ln in raw_lines:
        t = ln.strip()
        if t.startswith("<!--") and t.endswith("-->"):
            continue
        lines.append(ln)

    i = 0
    if lines:
        first = doc.add_paragraph(lines[0].strip())
        if first.runs:
            first.runs[0].bold = True
            first.runs[0].font.size = Pt(16)
        i = 1
    if i < len(lines):
        second = lines[i].strip()
        if second:
            p2 = doc.add_paragraph(second)
            if p2.runs:
                p2.runs[0].bold = True
                p2.runs[0].font.size = Pt(11)
        i += 1
    if i < len(lines):
        third = lines[i].strip()
        if third:
            doc.add_paragraph(third)
        i += 1

    if template_links:
        p = doc.add_paragraph("🔗 ")
        add_hyperlink(p, "LinkedIn", template_links.get("linkedin", "https://www.linkedin.com/"))
        p.add_run(" | 💻 ")
        add_hyperlink(p, "GitHub", template_links.get("github", "https://github.com/"))
        p.add_run(" | 🌐 ")
        add_hyperlink(p, "Portfolio", template_links.get("portfolio", "https://"))
    elif profile:
        p = doc.add_paragraph("🔗 ")
        add_hyperlink(p, "LinkedIn", profile.get("linkedin", "https://www.linkedin.com/"))
        p.add_run(" | 💻 ")
        add_hyperlink(p, "GitHub", profile.get("github", "https://github.com/"))
        p.add_run(" | 🌐 ")
        add_hyperlink(p, "Portfolio", profile.get("portfolio", "https://"))

    for raw in lines[i:]:
        line = raw.rstrip()
        if not line.strip():
            doc.add_paragraph("")
            continue

        clean = line.strip()
        low = clean.lower()
        if "linkedin" in low and "github" in low and "portfolio" in low:
            # We already render a clickable hyperlink row above.
            continue
        upper = clean.upper()
        if upper in section_headers:
            p = doc.add_paragraph(clean)
            if p.runs:
                p.runs[0].bold = True
                p.runs[0].font.size = Pt(11.5)
            continue

        if clean.startswith("### ") or clean.startswith("## ") or clean.startswith("# "):
            title = clean.lstrip("# ").strip()
            p = doc.add_paragraph(title)
            if p.runs:
                p.runs[0].bold = True
            continue

        if clean.startswith("- "):
            p = doc.add_paragraph(clean[2:].strip(), style="List Bullet")
            continue

        # Emphasize likely role/company lines for readability.
        if "|" in clean and ("–" in clean or "-" in clean):
            p = doc.add_paragraph(clean)
            if p.runs:
                p.runs[0].bold = True
            continue

        doc.add_paragraph(clean)

    return save_doc_with_fallback(doc, path)


def write_cover_letter_docx(text: str, path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    style_doc_defaults(doc)
    for raw in text.splitlines():
        line = raw.rstrip()
        if not line.strip():
            doc.add_paragraph("")
            continue
        if line.startswith("### "):
            p = doc.add_paragraph(line[4:].strip())
            p.style = "Heading 3"
            continue
        if line.startswith("## "):
            p = doc.add_paragraph(line[3:].strip())
            p.style = "Heading 2"
            continue
        if line.startswith("# "):
            p = doc.add_paragraph(line[2:].strip())
            p.style = "Heading 1"
            continue
        if line.startswith("- "):
            p = doc.add_paragraph(line[2:].strip())
            try:
                p.style = "List Bullet"
            except Exception:
                pass
            continue
        doc.add_paragraph(line)
    return save_doc_with_fallback(doc, path)


def save_doc_with_fallback(doc: Document, path: Path) -> Path:
    try:
        doc.save(path)
        return path
    except PermissionError:
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        alt = path.with_name(f"{path.stem}_{ts}{path.suffix}")
        doc.save(alt)
        return alt


def resolve_edge_executable() -> str:
    found = shutil.which("msedge")
    if found:
        return found
    candidates = [
        r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
        r"C:\Program Files\Microsoft\Edge\Application\msedge.exe",
    ]
    for path in candidates:
        if Path(path).exists():
            return path
    raise FileNotFoundError("Microsoft Edge executable not found.")


def ensure_system_edge_cdp(cdp_url: str, start_url: str = "about:blank") -> None:
    parsed = urlparse(cdp_url)
    port = parsed.port or 9222
    version_url = f"http://127.0.0.1:{port}/json/version"
    def _is_ready() -> bool:
        try:
            with httpx.Client(timeout=2.5) as client:
                resp = client.get(version_url)
                return resp.status_code == 200
        except Exception:
            return False

    if _is_ready():
        return
    edge_exe = resolve_edge_executable()
    subprocess.Popen(
        [
            edge_exe,
            f"--remote-debugging-port={port}",
            "--new-window",
            start_url,
        ],
        shell=False,
    )
    deadline = time.time() + 15.0
    while time.time() < deadline:
        if _is_ready():
            return
        time.sleep(0.5)


def write_csv(rows: List[dict], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    prepared_rows = prepare_export_rows(rows)
    focus_headers = resolve_export_headers(prepared_rows, AUDIT_FOCUS_HEADERS, AUDIT_ALWAYS_KEEP_HEADERS)
    detail_headers = resolve_export_headers(prepared_rows, AUDIT_DETAIL_HEADERS)
    base_headers = focus_headers + detail_headers
    extra_headers = build_extra_headers(prepared_rows, base_headers)
    headers = base_headers + extra_headers
    try:
        with path.open("w", encoding="utf-8-sig", newline="") as f:
            writer = csv.DictWriter(f, fieldnames=headers)
            writer.writeheader()
            for prepared in prepared_rows:
                writer.writerow({header: prepared.get(header, "") for header in headers})
    except PermissionError:
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        alt = path.with_name(f"{path.stem}_{ts}{path.suffix}")
        with alt.open("w", encoding="utf-8-sig", newline="") as f:
            writer = csv.DictWriter(f, fieldnames=headers)
            writer.writeheader()
            for prepared in prepared_rows:
                writer.writerow({header: prepared.get(header, "") for header in headers})
        print(f"Warning: tracker CSV is locked. Wrote fallback file: {alt}")


def add_dashboard_sheet(wb: Workbook, rows: List[dict]) -> None:
    if "dashboard" in wb.sheetnames:
        del wb["dashboard"]
    ws = wb.create_sheet("dashboard")

    ws["A1"] = "Metric"
    ws["B1"] = "Value"

    total = len(rows)
    scores = [parse_numeric_score(r.get("match_score", "")) for r in rows if parse_numeric_score(r.get("match_score", "")) > 0]
    avg_score = round(sum(scores) / len(scores), 1) if scores else 0.0
    high_match = sum(1 for r in rows if parse_numeric_score(r.get("match_score", "")) >= 70)
    posted_known = 0
    fresh_7d = 0
    days_values: List[int] = []
    for r in rows:
        posted = normalize_posted_date(r.get("job_posted_date", ""))
        if posted:
            posted_known += 1
            d = days_since_posted(posted)
            if d.isdigit():
                dv = int(d)
                days_values.append(dv)
                if dv <= 7:
                    fresh_7d += 1

    metric_rows = [
        ("Total Roles", total),
        ("Average Match Score", avg_score),
        ("High Match (>=70)", high_match),
        ("High Match Rate %", round((high_match / total) * 100, 1) if total else 0.0),
        ("Posted-Date Coverage %", round((posted_known / total) * 100, 1) if total else 0.0),
        ("Fresh Jobs (<=7d) %", round((fresh_7d / total) * 100, 1) if total else 0.0),
        ("Avg Days Since Posted", round(sum(days_values) / len(days_values), 1) if days_values else 0.0),
    ]
    for i, (k, v) in enumerate(metric_rows, start=2):
        ws.cell(row=i, column=1, value=k)
        ws.cell(row=i, column=2, value=v)

    # Status distribution table.
    ws["D1"] = "Status"
    ws["E1"] = "Count"
    status_counts: Dict[str, int] = {}
    for r in rows:
        status = normalize_status_value(r.get("status", ""))
        status_counts[status] = status_counts.get(status, 0) + 1
    status_items = sorted(status_counts.items(), key=lambda x: (-x[1], x[0]))
    for i, (status, count) in enumerate(status_items, start=2):
        ws.cell(row=i, column=4, value=status)
        ws.cell(row=i, column=5, value=count)

    # Platform performance table.
    ws["G1"] = "Platform"
    ws["H1"] = "Avg Match"
    ws["I1"] = "Count"
    platform_buckets: Dict[str, List[float]] = {}
    for r in rows:
        p = str(r.get("platform", "")).strip().lower() or "unknown"
        s = parse_numeric_score(r.get("match_score", ""))
        if s > 0:
            platform_buckets.setdefault(p, []).append(s)
        else:
            platform_buckets.setdefault(p, [])
    platform_items = sorted(platform_buckets.items(), key=lambda x: x[0])
    for i, (platform, vals) in enumerate(platform_items, start=2):
        ws.cell(row=i, column=7, value=platform)
        ws.cell(row=i, column=8, value=round(sum(vals) / len(vals), 1) if vals else 0.0)
        ws.cell(row=i, column=9, value=len(vals))

    # Posting freshness buckets table.
    ws["K1"] = "Posted Bucket"
    ws["L1"] = "Count"
    buckets = {"0-3 days": 0, "4-7 days": 0, "8-14 days": 0, "15+ days": 0, "Unknown": 0}
    for r in rows:
        d = days_since_posted(r.get("job_posted_date", ""))
        if not d.isdigit():
            buckets["Unknown"] += 1
            continue
        v = int(d)
        if v <= 3:
            buckets["0-3 days"] += 1
        elif v <= 7:
            buckets["4-7 days"] += 1
        elif v <= 14:
            buckets["8-14 days"] += 1
        else:
            buckets["15+ days"] += 1
    bucket_items = list(buckets.items())
    for i, (bucket, count) in enumerate(bucket_items, start=2):
        ws.cell(row=i, column=11, value=bucket)
        ws.cell(row=i, column=12, value=count)

    # Top missing keywords table.
    ws["N1"] = "Missing Keyword"
    ws["O1"] = "Count"
    miss_freq: Dict[str, int] = {}
    for r in rows:
        raw_missing = str(r.get("missing_keywords", "")).strip()
        if not raw_missing or raw_missing.lower() == "none":
            continue
        for kw in [x.strip() for x in raw_missing.split(",") if x.strip()]:
            if keyword_is_nonessential(kw):
                continue
            key = normalize_keyword_key(kw)
            if not key:
                continue
            miss_freq[key] = miss_freq.get(key, 0) + 1
    top_missing = sorted(miss_freq.items(), key=lambda x: (-x[1], x[0]))[:8]
    for i, (kw, count) in enumerate(top_missing, start=2):
        ws.cell(row=i, column=14, value=kw)
        ws.cell(row=i, column=15, value=count)

    if total == 0:
        ws["A10"] = "No rows in audit yet. Run scrape + run to populate charts."
        return

    # Chart 1: Status pie.
    if status_items:
        pie = PieChart()
        pie.title = "Status Distribution"
        data = Reference(ws, min_col=5, min_row=1, max_row=1 + len(status_items))
        labels = Reference(ws, min_col=4, min_row=2, max_row=1 + len(status_items))
        pie.add_data(data, titles_from_data=True)
        pie.set_categories(labels)
        pie.height = 7
        pie.width = 9
        ws.add_chart(pie, "A12")

    # Chart 2: Platform average match.
    if platform_items:
        bar = BarChart()
        bar.type = "col"
        bar.style = 10
        bar.title = "Average Match by Platform"
        bar.y_axis.title = "Score"
        bar.x_axis.title = "Platform"
        data = Reference(ws, min_col=8, min_row=1, max_row=1 + len(platform_items))
        cats = Reference(ws, min_col=7, min_row=2, max_row=1 + len(platform_items))
        bar.add_data(data, titles_from_data=True)
        bar.set_categories(cats)
        bar.height = 7
        bar.width = 9
        ws.add_chart(bar, "F12")

    # Chart 3: Posting freshness buckets.
    line = LineChart()
    line.style = 13
    line.title = "Posting Freshness"
    line.y_axis.title = "Count"
    line.x_axis.title = "Days Since Posted"
    data = Reference(ws, min_col=12, min_row=1, max_row=1 + len(bucket_items))
    cats = Reference(ws, min_col=11, min_row=2, max_row=1 + len(bucket_items))
    line.add_data(data, titles_from_data=True)
    line.set_categories(cats)
    line.height = 7
    line.width = 9
    ws.add_chart(line, "K12")

    # Chart 4: Top missing keywords.
    if top_missing:
        miss_bar = BarChart()
        miss_bar.type = "bar"
        miss_bar.style = 11
        miss_bar.title = "Top Missing Keywords"
        miss_bar.x_axis.title = "Count"
        miss_bar.y_axis.title = "Keyword"
        data = Reference(ws, min_col=15, min_row=1, max_row=1 + len(top_missing))
        cats = Reference(ws, min_col=14, min_row=2, max_row=1 + len(top_missing))
        miss_bar.add_data(data, titles_from_data=True)
        miss_bar.set_categories(cats)
        miss_bar.height = 7
        miss_bar.width = 9
        ws.add_chart(miss_bar, "P12")


def write_xlsx(rows: List[dict], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "priority_view"

    prepared_rows = prepare_export_rows(rows)
    priority_headers = resolve_export_headers(prepared_rows, AUDIT_FOCUS_HEADERS, AUDIT_ALWAYS_KEEP_HEADERS)
    ws.append(priority_headers)
    for prepared in prepared_rows:
        ws.append([prepared.get(h, "") for h in priority_headers])

    raw_ws = wb.create_sheet("raw_details")
    detail_headers = resolve_export_headers(prepared_rows, AUDIT_DETAIL_HEADERS)
    base_headers = priority_headers + detail_headers
    extra_headers = build_extra_headers(prepared_rows, base_headers)
    headers = base_headers + extra_headers
    raw_ws.append(headers)
    for prepared in prepared_rows:
        raw_ws.append([prepared.get(h, "") for h in headers])
    add_dashboard_sheet(wb, prepared_rows)
    try:
        wb.save(path)
    except PermissionError:
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        alt = path.with_name(f"{path.stem}_{ts}{path.suffix}")
        wb.save(alt)
        print(f"Warning: xlsx is locked. Wrote fallback file: {alt}")


def read_xlsx_rows(path: Path) -> List[dict]:
    if not path.exists():
        return []
    wb = load_workbook(path)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(h) if h is not None else "" for h in rows[0]]
    data: List[dict] = []
    for r in rows[1:]:
        row_dict = {}
        for i, h in enumerate(headers):
            if not h:
                continue
            v = r[i] if i < len(r) else ""
            row_dict[h] = "" if v is None else str(v)
        if row_dict:
            data.append(row_dict)
    return data


def assist_apply(platform: str, limit: int = 20, headless: bool = False) -> None:
    ensure_playwright_ready()
    if not AUDIT_XLSX_FILE.exists():
        raise FileNotFoundError(f"Audit file not found: {AUDIT_XLSX_FILE}. Run pipeline first.")

    rows = read_xlsx_rows(AUDIT_XLSX_FILE)
    pending_pairs: List[Tuple[int, dict]] = []
    for idx, row in enumerate(rows):
        if row.get("platform", "").lower() == platform.lower() and row.get("status", "") == "ready_to_apply":
            pending_pairs.append((idx, row))
    pending_pairs = pending_pairs[:limit]

    if not pending_pairs:
        print(f"No pending rows for platform={platform}")
        return

    profile_path = PROFILE_DIR / platform.lower()
    if not profile_path.exists():
        raise FileNotFoundError(
            f"Missing browser profile for {platform}. Run auth first: "
            f"py scripts/job_workflow.py auth --platform {platform}"
        )

    with sync_playwright() as p:
        context = p.chromium.launch_persistent_context(
            user_data_dir=str(profile_path),
            headless=headless,
            viewport={"width": 1440, "height": 900},
        )
        page = context.new_page()
        for idx, row in pending_pairs:
            url = str(row.get("job_url", "")).strip()
            if not url:
                rows[idx]["status"] = "skipped_no_url"
                continue

            print("=" * 60)
            print(f"Opening: {row.get('title', '')} @ {row.get('company', '')}")
            print(f"Resume: {row.get('resume_file', '')}")
            try:
                page.goto(url, wait_until="domcontentloaded")
            except PlaywrightTimeoutError:
                pass
            print("Please complete application manually in browser.")
            print("Type one of: submitted / skipped / failed")
            decision = input("> ").strip().lower() or "skipped"
            if decision not in {"submitted", "skipped", "failed"}:
                decision = "skipped"
            rows[idx]["status"] = decision
            rows[idx]["applied_at"] = datetime.now().isoformat(timespec="seconds")
        context.close()

    sync_tracker_outputs(rows)
    print(f"Audit updated: {AUDIT_XLSX_FILE}")


def click_first(page, selectors: List[str]) -> bool:
    scopes = [page] + list(page.frames)
    for scope in scopes:
        for sel in selectors:
            try:
                loc = scope.locator(sel).first
                if loc.count() > 0 and loc.is_visible(timeout=1200):
                    loc.click(timeout=2200)
                    return True
            except Exception:
                continue
    return False


def try_auto_click_apply(page, platform: str) -> str:
    platform = platform.lower()
    if platform == "seek":
        selectors = [
            "button:has-text('Apply')",
            "button:has-text('Quick apply')",
            "a:has-text('Apply')",
            "[data-automation*='apply']",
        ]
    else:
        selectors = [
            "button:has-text('Easy Apply')",
            "button:has-text('Apply')",
            "a:has-text('Easy Apply')",
            "[aria-label*='Easy Apply']",
        ]
    ok = click_first(page, selectors)
    return "auto_clicked_apply" if ok else "auto_click_not_found"


def click_apply_chain(page) -> List[str]:
    actions: List[str] = []
    selectors = [
        "button:has-text('Apply')",
        "button:has-text('Apply now')",
        "button:has-text('Quick apply')",
        "button:has-text('Easy Apply')",
        "a:has-text('Apply')",
        "a:has-text('Apply now')",
        "a:has-text('Apply on company site')",
        "a:has-text('Apply on employer site')",
        "a:has-text('Continue')",
        "button:has-text('Continue')",
    ]
    for _ in range(4):
        clicked = click_first(page, selectors)
        if not clicked:
            break
        actions.append("clicked_apply_step")
        page.wait_for_timeout(1200)
    return actions


def fill_first(page, selectors: List[str], value: str) -> bool:
    if not value:
        return False
    scopes = [page] + list(page.frames)
    for scope in scopes:
        for sel in selectors:
            try:
                loc = scope.locator(sel).first
                if loc.count() > 0 and loc.is_visible(timeout=900):
                    existing = ""
                    try:
                        existing = (loc.input_value(timeout=900) or "").strip()
                    except Exception:
                        existing = ""
                    # Strict safety: do not overwrite existing user/provided content.
                    if existing:
                        continue
                    loc.fill(value, timeout=2200)
                    return True
            except Exception:
                continue
    return False


def autofill_form(page, profile: dict, resume_file: str, cover_letter_text: str) -> Dict[str, int]:
    # Strict whitelist-only autofill to avoid hallucinated form mappings.
    stats = {"fields": 0, "uploads": 0, "coverletter": 0}
    mappings = {
        profile.get("name", "").split(" ")[0]: [
            "input[name*='first']",
            "input[id*='first']",
            "input[placeholder*='First']",
        ],
        " ".join(profile.get("name", "").split(" ")[1:]): [
            "input[name*='last']",
            "input[id*='last']",
            "input[placeholder*='Last']",
        ],
    }
    if fill_first(page, mappings.get(profile.get("name", "").split(" ")[0], []), profile.get("name", "").split(" ")[0]):
        stats["fields"] += 1
    last_name = " ".join(profile.get("name", "").split(" ")[1:]).strip()
    if last_name and fill_first(page, mappings.get(last_name, []), last_name):
        stats["fields"] += 1

    field_map = [
        (profile.get("email", ""), ["input[type='email']", "input[name*='email']", "input[id*='email']"]),
        (profile.get("phone", ""), ["input[type='tel']", "input[name*='phone']", "input[id*='phone']"]),
        (profile.get("address_line", ""), ["input[name*='address']", "input[id*='address']", "input[placeholder*='Address']"]),
        (profile.get("location", ""), ["input[name*='city']", "input[id*='city']", "input[placeholder*='City']"]),
    ]
    for value, selectors in field_map:
        if fill_first(page, selectors, value):
            stats["fields"] += 1

    # Select-based demographic fields where applicable.
    for value, selectors in [
        (profile.get("gender", ""), ["select[name*='gender']", "select[id*='gender']"]),
        (profile.get("has_disability", ""), ["select[name*='disab']", "select[id*='disab']"]),
        (profile.get("is_lgbti", ""), ["select[name*='lgbt']", "select[id*='lgbt']"]),
    ]:
        for sel in selectors:
            try:
                loc = page.locator(sel).first
                if loc.count() > 0:
                    ok = False
                    try:
                        loc.select_option(label=value, timeout=1500)
                        ok = True
                    except Exception:
                        try:
                            loc.select_option(value=value.lower(), timeout=1500)
                            ok = True
                        except Exception:
                            ok = False
                    if ok:
                        stats["fields"] += 1
                        break
            except Exception:
                continue

    # Upload CV
    try:
        file_input = page.locator("input[type='file']").first
        if file_input.count() > 0:
            file_input.set_input_files(resume_file, timeout=3000)
            stats["uploads"] += 1
    except Exception:
        pass

    # Cover letter textarea
    if cover_letter_text:
        if fill_first(
            page,
            ["textarea[name*='cover']", "textarea[id*='cover']", "textarea[placeholder*='cover']", "textarea[name*='letter']", "textarea[id*='letter']"],
            cover_letter_text[:3500],
        ):
            stats["coverletter"] += 1

    return stats


def read_cover_letter_text(path: str) -> str:
    p = Path(path)
    if not p.exists():
        return ""
    try:
        doc = Document(str(p))
        return "\n".join([x.text for x in doc.paragraphs if x.text.strip()])
    except Exception:
        return ""


def scrape_seek_jobs(
    query: str,
    location: str,
    limit: int = 20,
    channel: str = "msedge",
    headless: bool = False,
    use_system_edge: bool = False,
    cdp_url: str = "http://127.0.0.1:9222",
) -> List[dict]:
    ensure_playwright_ready()
    debug_seek = os.getenv("SEEK_DEBUG", "").strip().lower() in {"1", "true", "yes", "on"}
    list_timeout_ms = int(os.getenv("SEEK_LIST_TIMEOUT_MS", "25000"))
    detail_timeout_ms = int(os.getenv("SEEK_DETAIL_TIMEOUT_MS", "12000"))
    max_candidates = int(os.getenv("SEEK_MAX_CANDIDATES", str(max(limit * 2, 12))))
    max_detail_attempts = int(os.getenv("SEEK_MAX_DETAIL_ATTEMPTS", "1"))
    max_runtime_s = int(os.getenv("SEEK_MAX_RUNTIME_S", "60"))
    deadline_ts = time.time() + max_runtime_s

    def seek_log(message: str) -> None:
        if debug_seek:
            print(f"[seek-debug] {message}")

    def to_seek_slug(value: str) -> str:
        clean = re.sub(r"[^A-Za-z0-9 ]+", " ", (value or "").strip())
        clean = re.sub(r"\s+", "-", clean).strip("-")
        return clean or "jobs"

    def norm_txt(value: str) -> str:
        return re.sub(r"\s+", " ", (value or "")).strip()

    def extract_jobposting_from_ldjson(raw: str) -> Optional[dict]:
        try:
            payload = json.loads(raw)
        except Exception:
            return None
        stack = [payload]
        while stack:
            item = stack.pop()
            if isinstance(item, list):
                stack.extend(item)
                continue
            if not isinstance(item, dict):
                continue
            t = str(item.get("@type", "")).lower()
            if "jobposting" in t:
                return item
            for k in ("@graph", "graph", "itemListElement"):
                v = item.get(k)
                if isinstance(v, (list, dict)):
                    stack.append(v)
        return None

    def parse_seek_detail(page2, fallback_title: str, fallback_loc: str) -> dict:
        body = ""
        try:
            body = page2.inner_text("body")
        except Exception:
            body = ""
        title = ""
        company = ""
        loc2 = ""
        jd_text = body or ""
        posted_date = ""
        try:
            title = norm_txt(page2.locator("h1").first.inner_text(timeout=1000))
        except Exception:
            title = ""
        try:
            company = norm_txt(
                page2.locator("a[data-automation='advertiser-name'], [data-automation='advertiser-name']").first.inner_text(timeout=1000)
            )
        except Exception:
            company = ""
        try:
            loc2 = norm_txt(page2.locator("[data-automation='job-detail-location']").first.inner_text(timeout=1000))
        except Exception:
            loc2 = ""
        try:
            jd = page2.locator("[data-automation='jobAdDetails']").first.inner_text(timeout=1200)
            if jd and len(jd.strip()) > 100:
                jd_text = jd
        except Exception:
            pass

        ld_scripts = []
        try:
            ld_scripts = page2.locator("script[type='application/ld+json']").all_text_contents()
        except Exception:
            ld_scripts = []
        for raw in ld_scripts:
            jobposting = extract_jobposting_from_ldjson(raw)
            if not jobposting:
                continue
            title = title or norm_txt(str(jobposting.get("title", "")))
            org = jobposting.get("hiringOrganization") or {}
            if isinstance(org, dict):
                company = company or norm_txt(str(org.get("name", "")))
            desc = jobposting.get("description")
            if isinstance(desc, str) and len(desc.strip()) > 50:
                desc_clean = norm_txt(re.sub(r"<[^>]+>", " ", html.unescape(desc)))
                if len(desc_clean) > 100:
                    jd_text = desc_clean
            job_loc = jobposting.get("jobLocation")
            if isinstance(job_loc, dict):
                job_loc = [job_loc]
            if isinstance(job_loc, list) and job_loc:
                addr = (job_loc[0] or {}).get("address", {})
                if isinstance(addr, dict):
                    loc_bits = [
                        str(addr.get("addressLocality", "")).strip(),
                        str(addr.get("addressRegion", "")).strip(),
                    ]
                    loc_joined = " ".join([x for x in loc_bits if x]).strip()
                    if loc_joined:
                        loc2 = loc2 or loc_joined
            posted_date = normalize_posted_date(jobposting.get("datePosted", "") or jobposting.get("validThrough", ""))
            break

        # Seek-specific posted-date metadata fallback.
        if not posted_date:
            date_snippets: List[str] = []
            for sel in [
                "[data-automation='job-detail-date']",
                "[data-automation='jobListingDate']",
                "[data-automation='job-detail-summary']",
                "[data-automation='jobDetailTop']",
                "time",
            ]:
                try:
                    parts = page2.locator(sel).all_inner_texts()
                    if parts:
                        date_snippets.extend([norm_txt(x) for x in parts if norm_txt(x)])
                except Exception:
                    continue
            try:
                datetime_attrs = page2.locator("time").evaluate_all("els => els.map(e => e.getAttribute('datetime') || '')")
                if datetime_attrs:
                    date_snippets.extend([norm_txt(x) for x in datetime_attrs if norm_txt(x)])
            except Exception:
                pass
            joined = " | ".join(date_snippets)
            posted_date = infer_posted_date_from_text(joined)

        if not loc2:
            mloc = re.search(r"([A-Za-z ]+,?\s+QLD|Brisbane QLD)", body or "", flags=re.I)
            if mloc:
                loc2 = norm_txt(mloc.group(1))
        if not posted_date:
            posted_date = infer_posted_date_from_text(body)
        quick = ("quick apply" in (body or "").lower()) or ("easy apply" in (body or "").lower())
        application_url = page2.url
        for sel in [
            "a:has-text('Apply on company site')",
            "a:has-text('Apply on employer site')",
            "a:has-text('Apply now')",
            "button:has-text('Apply now')",
        ]:
            try:
                locator = page2.locator(sel).first
                href = locator.get_attribute("href")
                if href:
                    if href.startswith("/"):
                        href = "https://www.seek.com.au" + href
                    application_url = href
                    break
            except Exception:
                continue
        return {
            "title": title or norm_txt(fallback_title) or "Unknown title",
            "company": company or "Unknown company",
            "location": loc2 or fallback_loc,
            "jd_text": (jd_text or "")[:12000],
            "job_posted_date": posted_date,
            "application_url": application_url,
            "quick_apply": quick,
            "notes": "quick_apply_detected" if quick else "standard_apply",
        }

    def parse_seek_detail_html(html_text: str, job_url: str, fallback_title: str, fallback_loc: str) -> dict:
        soup = BeautifulSoup(html_text or "", "html.parser")
        page_text = norm_txt(soup.get_text(" ", strip=True))
        title = ""
        company = ""
        loc2 = ""
        jd_text = page_text
        posted_date = ""
        h1 = soup.select_one("h1")
        if h1:
            title = norm_txt(h1.get_text(" ", strip=True))
        company_node = soup.select_one("a[data-automation='advertiser-name'], [data-automation='advertiser-name']")
        if company_node:
            company = norm_txt(company_node.get_text(" ", strip=True))
        loc_node = soup.select_one("[data-automation='job-detail-location']")
        if loc_node:
            loc2 = norm_txt(loc_node.get_text(" ", strip=True))
        jd_node = soup.select_one("[data-automation='jobAdDetails']")
        if jd_node:
            jd_candidate = norm_txt(jd_node.get_text(" ", strip=True))
            if len(jd_candidate) > 100:
                jd_text = jd_candidate

        for script in soup.select("script[type='application/ld+json']"):
            raw = script.string or script.get_text() or ""
            jobposting = extract_jobposting_from_ldjson(raw)
            if not jobposting:
                continue
            title = title or norm_txt(str(jobposting.get("title", "")))
            org = jobposting.get("hiringOrganization") or {}
            if isinstance(org, dict):
                company = company or norm_txt(str(org.get("name", "")))
            desc = jobposting.get("description")
            if isinstance(desc, str) and len(desc.strip()) > 50:
                desc_clean = norm_txt(re.sub(r"<[^>]+>", " ", html.unescape(desc)))
                if len(desc_clean) > 100:
                    jd_text = desc_clean
            job_loc = jobposting.get("jobLocation")
            if isinstance(job_loc, dict):
                job_loc = [job_loc]
            if isinstance(job_loc, list) and job_loc:
                addr = (job_loc[0] or {}).get("address", {})
                if isinstance(addr, dict):
                    loc_bits = [
                        str(addr.get("addressLocality", "")).strip(),
                        str(addr.get("addressRegion", "")).strip(),
                    ]
                    loc_joined = " ".join([x for x in loc_bits if x]).strip()
                    if loc_joined:
                        loc2 = loc2 or loc_joined
            posted_date = normalize_posted_date(jobposting.get("datePosted", "") or jobposting.get("validThrough", ""))
            break

        if not posted_date:
            posted_date = infer_posted_date_from_text(page_text)

        quick = ("quick apply" in page_text.lower()) or ("easy apply" in page_text.lower())
        application_url = job_url
        for a in soup.select("a[href]"):
            text = norm_txt(a.get_text(" ", strip=True)).lower()
            if text in {"apply now", "apply on company site", "apply on employer site"}:
                href = (a.get("href") or "").strip()
                if href:
                    if href.startswith("/"):
                        href = "https://www.seek.com.au" + href
                    application_url = href
                    break
        return {
            "title": title or norm_txt(fallback_title) or "Unknown title",
            "company": company or "Unknown company",
            "location": loc2 or fallback_loc,
            "jd_text": (jd_text or "")[:12000],
            "job_posted_date": posted_date,
            "application_url": application_url,
            "quick_apply": quick,
            "notes": "quick_apply_detected" if quick else "standard_apply",
        }

    def scrape_seek_jobs_http(search_url: str, fallback_loc: str, max_jobs: int) -> List[dict]:
        headers = {
            "User-Agent": (
                "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                "AppleWebKit/537.36 (KHTML, like Gecko) "
                "Chrome/125.0 Safari/537.36"
            )
        }
        out: List[dict] = []
        candidates: List[Tuple[str, str]] = []
        seen_urls: Set[str] = set()
        with httpx.Client(timeout=25, follow_redirects=True, headers=headers) as client:
            resp = client.get(search_url)
            resp.raise_for_status()
            soup = BeautifulSoup(resp.text, "html.parser")
            for a in soup.select("a[href*='/job/']"):
                href = (a.get("href") or "").strip()
                if not href or "/job/" not in href:
                    continue
                if href.startswith("/"):
                    href = "https://www.seek.com.au" + href
                href = href.split("?")[0]
                if href in seen_urls:
                    continue
                seen_urls.add(href)
                title = norm_txt(a.get_text(" ", strip=True))
                candidates.append((href, title))
                if len(candidates) >= max(max_jobs * 4, 40):
                    break
            seek_log(f"http_candidates={len(candidates)}")
            for href, title in candidates:
                try:
                    r = client.get(href)
                    if r.status_code >= 400:
                        continue
                    parsed = parse_seek_detail_html(r.text, href, title, fallback_loc)
                    out.append(
                        {
                            "job_url": href,
                            "application_url": parsed.get("application_url", href),
                            "title": parsed["title"],
                            "company": norm_txt(parsed["company"]),
                            "location": norm_txt(parsed["location"]),
                            "job_posted_date": parsed.get("job_posted_date", ""),
                            "jd_text": parsed["jd_text"],
                            "notes": f"{parsed['notes']}; http_fallback",
                            "quick_apply": bool(parsed["quick_apply"]),
                        }
                    )
                except Exception:
                    continue
                if len(out) >= max_jobs:
                    break
        return out

    q = (query or "").strip() or "business graduate"
    loc = (location or "").strip() or "Brisbane QLD"
    url = f"https://www.seek.com.au/{to_seek_slug(q)}-jobs/in-{to_seek_slug(loc)}"
    jobs: List[dict] = []
    http_first = os.getenv("SEEK_HTTP_FIRST", "1").strip().lower() not in {"0", "false", "no", "off"}
    if http_first:
        try:
            seek_log("http_first enabled")
            jobs = scrape_seek_jobs_http(url, loc, limit)
        except Exception as exc:
            seek_log(f"http_first failed error={exc}")
            jobs = []
        if jobs:
            seek_log(f"http_first_success jobs={len(jobs)}")
            jobs.sort(key=lambda x: (not x.get("quick_apply", False), x.get("title", "")))
            return jobs[:limit]

    def open_seek_session(pw, start_url: str) -> Tuple[object, object, bool]:
        browser = None
        context = None
        created_context = False
        state_path = storage_state_path("seek")
        if state_path.exists():
            browser, context, mode = open_saved_platform_context(
                pw,
                "seek",
                headless=headless,
                channel=channel,
                prefer_headless_saved_state=True,
            )
            return browser, context, mode != "persistent_profile"
        if use_system_edge:
            try:
                ensure_system_edge_cdp(cdp_url, start_url=start_url)
                browser = pw.chromium.connect_over_cdp(cdp_url)
                context = browser.contexts[0] if browser.contexts else browser.new_context()
                created_context = len(browser.contexts) == 0
                return browser, context, created_context
            except Exception:
                print(
                    "Warning: failed to attach system Edge CDP; falling back to saved Seek session/profile."
                )
        browser, context, mode = open_saved_platform_context(
            pw,
            "seek",
            headless=headless,
            channel=channel,
            prefer_headless_saved_state=True,
        )
        return browser, context, mode != "persistent_profile"

    def close_seek_session(browser, context, created_context: bool) -> None:
        try:
            if context is not None and (not use_system_edge or created_context):
                context.close()
        except Exception:
            pass
        try:
            if browser is not None and (not use_system_edge or created_context):
                browser.close()
        except Exception:
            pass

    with sync_playwright() as p:
        if time.time() > deadline_ts:
            return jobs[:limit]
        browser, context, created_context = open_seek_session(p, url)
        page = None
        detail_page = None
        candidates: List[Tuple[str, str]] = []
        try:
            if time.time() > deadline_ts:
                return jobs[:limit]
            seek_log(f"open_list url={url}")
            page = context.new_page()
            page.goto(url, wait_until="commit", timeout=list_timeout_ms)
            page.wait_for_timeout(1800)
            page.mouse.wheel(0, 2800)
            page.wait_for_timeout(900)
            cards = page.locator("a[href*='/job/']").all()[: max(limit * 8, 80)]
            seek_log(f"list_cards={len(cards)}")
            seen = set()
            for a in cards:
                try:
                    href = a.get_attribute("href") or ""
                    title = norm_txt(a.inner_text() or "")
                except Exception:
                    continue
                if not href or "/job/" not in href:
                    continue
                if href.startswith("/"):
                    href = "https://www.seek.com.au" + href
                href = href.split("?")[0]
                if href in seen:
                    continue
                seen.add(href)
                candidates.append((href, title))
                if len(candidates) >= max_candidates:
                    break
            if not candidates:
                try:
                    html_content = page.content()
                    for href in re.findall(r"https://www\.seek\.com\.au/job/\d+", html_content or ""):
                        if href in seen:
                            continue
                        seen.add(href)
                        candidates.append((href, ""))
                        if len(candidates) >= max_candidates:
                            break
                except Exception:
                    pass
            seek_log(f"candidate_count={len(candidates)}")
        finally:
            try:
                if page is not None and not page.is_closed():
                    page.close()
            except Exception:
                pass

        for href, title in candidates:
            if time.time() > deadline_ts:
                seek_log("deadline_reached_before_detail_loop")
                break
            attempt = 0
            while attempt < max_detail_attempts:
                if time.time() > deadline_ts:
                    seek_log("deadline_reached_inside_detail_attempts")
                    break
                attempt += 1
                try:
                    if context is None:
                        seek_log(f"reopen_session_for={href}")
                        browser, context, created_context = open_seek_session(p, href)
                    if detail_page is None or detail_page.is_closed():
                        detail_page = context.new_page()
                    seek_log(f"open_detail attempt={attempt} url={href}")
                    detail_page.goto(href, wait_until="commit", timeout=detail_timeout_ms)
                    detail_page.wait_for_timeout(700)
                    parsed = parse_seek_detail(detail_page, fallback_title=title, fallback_loc=loc)
                    jobs.append(
                        {
                            "job_url": href,
                            "application_url": parsed.get("application_url", href),
                            "title": parsed["title"],
                            "company": norm_txt(parsed["company"]),
                            "location": norm_txt(parsed["location"]),
                            "job_posted_date": parsed.get("job_posted_date", ""),
                            "jd_text": parsed["jd_text"],
                            "notes": parsed["notes"],
                            "quick_apply": bool(parsed["quick_apply"]),
                        }
                    )
                    seek_log(
                        "detail_ok title={title} company={company} posted={posted}".format(
                            title=parsed.get("title", "")[:80],
                            company=norm_txt(parsed.get("company", ""))[:60],
                            posted=parsed.get("job_posted_date", ""),
                        )
                    )
                    try:
                        if detail_page and not detail_page.is_closed():
                            detail_page.close()
                    except Exception:
                        pass
                    detail_page = None
                    break
                except (PlaywrightError, PlaywrightTimeoutError) as exc:
                    seek_log(f"detail_fail attempt={attempt} url={href} error={exc}")
                    print(f"Seek detail scrape warning for {href}: {exc}")
                    try:
                        if detail_page and not detail_page.is_closed():
                            detail_page.close()
                    except Exception:
                        pass
                    detail_page = None
                    close_seek_session(browser, context, created_context)
                    browser, context, created_context = None, None, False
                    if attempt >= max_detail_attempts:
                        break
                    continue
                except Exception:
                    seek_log(f"detail_fail_unknown attempt={attempt} url={href}")
                    try:
                        if detail_page and not detail_page.is_closed():
                            detail_page.close()
                    except Exception:
                        pass
                    detail_page = None
                    if attempt >= max_detail_attempts:
                        break
                    continue
            if len(jobs) >= limit:
                break

        try:
            if detail_page and not detail_page.is_closed():
                detail_page.close()
        except Exception:
            pass
        close_seek_session(browser, context, created_context)
    seek_log(f"scrape_done jobs={len(jobs)}")
    if not jobs:
        try:
            seek_log("playwright_empty_try_http_fallback")
            jobs = scrape_seek_jobs_http(url, loc, limit)
        except Exception as exc:
            seek_log(f"http_fallback_after_playwright failed error={exc}")
            jobs = []
    # Sort quick apply first for easier automation testing.
    jobs.sort(key=lambda x: (not x.get("quick_apply", False), x.get("title", "")))
    return jobs[:limit]


def scrape_linkedin_jobs(
    query: str,
    location: str,
    limit: int = 20,
    channel: str = "msedge",
    headless: bool = False,
    use_system_edge: bool = False,
    cdp_url: str = "http://127.0.0.1:9222",
) -> List[dict]:
    ensure_playwright_ready()
    def norm_txt(value: str) -> str:
        return re.sub(r"\s+", " ", (value or "")).strip()

    def dedupe_repeated_phrase(value: str) -> str:
        txt = norm_txt(value)
        if not txt:
            return txt
        parts = txt.split(" ")
        n = len(parts)
        if n >= 4 and n % 2 == 0:
            mid = n // 2
            if " ".join(parts[:mid]).lower() == " ".join(parts[mid:]).lower():
                return " ".join(parts[:mid]).strip()
        return txt

    def extract_jobposting_from_ldjson(raw: str) -> Optional[dict]:
        try:
            payload = json.loads(raw)
        except Exception:
            return None
        stack = [payload]
        while stack:
            item = stack.pop()
            if isinstance(item, list):
                stack.extend(item)
                continue
            if not isinstance(item, dict):
                continue
            t = str(item.get("@type", "")).lower()
            if "jobposting" in t:
                return item
            for k in ("@graph", "graph", "itemListElement"):
                v = item.get(k)
                if isinstance(v, (list, dict)):
                    stack.append(v)
        return None

    def parse_linkedin_detail(page2, fallback_title: str, fallback_company: str, fallback_loc: str) -> dict:
        title = ""
        company = ""
        loc2 = ""
        jd_text = ""
        posted_date = ""
        body = ""
        try:
            body = page2.inner_text("body")
        except Exception:
            body = ""
        for sel in ["h1.top-card-layout__title", "h1.t-24", "h1"]:
            try:
                title = norm_txt(page2.locator(sel).first.inner_text(timeout=900))
                if title:
                    break
            except Exception:
                continue
        for sel in [
            "a.topcard__org-name-link",
            ".topcard__flavor-row .topcard__flavor",
            "div.top-card-layout__card a",
            "span.jobs-unified-top-card__company-name",
        ]:
            try:
                company = norm_txt(page2.locator(sel).first.inner_text(timeout=900))
                if company:
                    break
            except Exception:
                continue
        # Fallback: parse from document title like "Role title | Company | LinkedIn".
        if not company:
            try:
                page_title = norm_txt(page2.title())
                chunks = [x.strip() for x in page_title.split("|") if x.strip()]
                if len(chunks) >= 2 and "linkedin" in chunks[-1].lower():
                    company = chunks[-2]
            except Exception:
                pass
        for sel in [
            "span.topcard__flavor--bullet",
            ".topcard__flavor.topcard__flavor--bullet",
            ".jobs-unified-top-card__bullet",
            "span.jobs-unified-top-card__workplace-type",
        ]:
            try:
                loc2 = norm_txt(page2.locator(sel).first.inner_text(timeout=900))
                if loc2:
                    break
            except Exception:
                continue
        for sel in [
            "div.show-more-less-html__markup",
            "div.jobs-description__content",
            "article.jobs-description",
            "section.description",
        ]:
            try:
                jd_text = norm_txt(page2.locator(sel).first.inner_text(timeout=1000))
                if jd_text:
                    break
            except Exception:
                continue

        ld_scripts = []
        try:
            ld_scripts = page2.locator("script[type='application/ld+json']").all_text_contents()
        except Exception:
            ld_scripts = []
        for raw in ld_scripts:
            jp = extract_jobposting_from_ldjson(raw)
            if not jp:
                continue
            title = title or norm_txt(str(jp.get("title", "")))
            org = jp.get("hiringOrganization") or {}
            if isinstance(org, dict):
                company = company or norm_txt(str(org.get("name", "")))
            desc = jp.get("description")
            if isinstance(desc, str) and len(desc.strip()) > 40:
                desc_clean = norm_txt(re.sub(r"<[^>]+>", " ", html.unescape(desc)))
                if len(desc_clean) > 100:
                    jd_text = jd_text or desc_clean
            posted_date = normalize_posted_date(jp.get("datePosted", "") or jp.get("validThrough", ""))
            break

        # LinkedIn-specific posted-date metadata fallback.
        if not posted_date:
            date_snippets: List[str] = []
            for sel in [
                "span.posted-time-ago__text",
                ".jobs-unified-top-card__primary-description-container",
                ".jobs-unified-top-card__subtitle-primary-grouping",
                ".topcard__flavor--metadata",
                "time",
            ]:
                try:
                    parts = page2.locator(sel).all_inner_texts()
                    if parts:
                        date_snippets.extend([norm_txt(x) for x in parts if norm_txt(x)])
                except Exception:
                    continue
            try:
                datetime_attrs = page2.locator("time").evaluate_all("els => els.map(e => e.getAttribute('datetime') || '')")
                if datetime_attrs:
                    date_snippets.extend([norm_txt(x) for x in datetime_attrs if norm_txt(x)])
            except Exception:
                pass
            joined = " | ".join(date_snippets)
            posted_date = infer_posted_date_from_text(joined)

        title = dedupe_repeated_phrase(title)
        if title:
            title = re.sub(r"\s+with verification$", "", title, flags=re.I).strip()
        easy = "easy apply" in (body or "").lower()
        application_url = page2.url
        for sel in [
            "a:has-text('Apply on company site')",
            "a:has-text('Apply')",
            "button:has-text('Easy Apply')",
            "a:has-text('Easy Apply')",
        ]:
            try:
                locator = page2.locator(sel).first
                href = locator.get_attribute("href")
                if href:
                    if href.startswith("/"):
                        href = urljoin("https://www.linkedin.com", href)
                    application_url = href
                    break
            except Exception:
                continue
        if not posted_date:
            posted_date = infer_posted_date_from_text(body)
        return {
            "title": title or fallback_title or "LinkedIn role",
            "company": company or fallback_company or "Unknown company",
            "location": loc2 or fallback_loc,
            "jd_text": (jd_text or body or "")[:12000],
            "job_posted_date": posted_date,
            "application_url": application_url,
            "quick_apply": easy,
            "notes": "easy_apply_detected" if easy else "linkedin_standard_apply",
        }

    def is_valid_linkedin_job(title: str, company: str, jd_text: str) -> bool:
        t = (title or "").strip().lower()
        c = (company or "").strip().lower()
        j = (jd_text or "").strip().lower()
        junk_markers = [
            "join linkedin",
            "sign in",
            "登录",
            "加入领英",
            "create your profile",
            "forgot password",
            "agree & join",
        ]
        if any(x in t for x in junk_markers):
            return False
        if any(x in j for x in ["同意并加入", "用户协议", "cookie 政策", "forgot password"]) and len(j) < 2500:
            return False
        if t in {"linkedin", "jobs", "home"}:
            return False
        if c in {"linkedin", "unknown company"} and len(t) < 4:
            return False
        return True

    q = (query or "").strip() or "business analyst"
    loc = (location or "").strip() or "Brisbane"
    url = f"https://www.linkedin.com/jobs/search/?keywords={quote_plus(q)}&location={quote_plus(loc)}"
    jobs: List[dict] = []
    with sync_playwright() as p:
        browser = None
        context = None
        created_context = False
        if use_system_edge and not storage_state_path("linkedin").exists():
            try:
                ensure_system_edge_cdp(cdp_url, start_url=url)
                browser = p.chromium.connect_over_cdp(cdp_url)
                context = browser.contexts[0] if browser.contexts else browser.new_context()
                created_context = len(browser.contexts) == 0
            except Exception:
                print(
                    "Warning: failed to attach system Edge CDP; falling back to saved LinkedIn session/profile."
                )
                browser, context, mode = open_saved_platform_context(
                    p,
                    "linkedin",
                    headless=headless,
                    channel=channel,
                    prefer_headless_saved_state=True,
                )
                created_context = mode != "persistent_profile"
                use_system_edge = False
        else:
            browser, context, mode = open_saved_platform_context(
                p,
                "linkedin",
                headless=headless,
                channel=channel,
                prefer_headless_saved_state=True,
            )
            created_context = mode != "persistent_profile"
            use_system_edge = False
        page = context.new_page()
        page.goto(url, wait_until="domcontentloaded", timeout=60000)
        page.wait_for_timeout(3200)
        page.mouse.wheel(0, 2400)
        page.wait_for_timeout(1200)
        links = page.locator("a[href*='/jobs/view/']").all()[: max(limit * 8, 80)]
        seen = set()
        for a in links:
            try:
                href = a.get_attribute("href") or ""
                title = dedupe_repeated_phrase(a.inner_text() or "")
            except Exception:
                continue
            if not href:
                continue
            href = href.split("?")[0]
            if href.startswith("/"):
                href = urljoin("https://www.linkedin.com", href)
            elif href.startswith("www."):
                href = "https://" + href
            if href in seen:
                continue
            seen.add(href)
            fallback_company = ""
            try:
                card = a.locator("xpath=ancestor::*[self::li or self::div][1]")
                if card.count() > 0:
                    fallback_company = norm_txt(
                        card.locator(
                            "h4, .base-search-card__subtitle, .job-search-card__subtitle-link, .artdeco-entity-lockup__subtitle"
                        ).first.inner_text(timeout=900)
                    )
            except Exception:
                fallback_company = ""
            parsed = {
                "title": title or "LinkedIn role",
                "company": fallback_company or "Unknown company",
                "location": loc,
                "jd_text": "",
                "notes": "linkedin_scraped_basic",
                "quick_apply": "easy apply" in (title or "").lower(),
            }
            try:
                page2 = context.new_page()
                page2.goto(href, wait_until="domcontentloaded", timeout=45000)
                page2.wait_for_timeout(1200)
                parsed = parse_linkedin_detail(page2, fallback_title=title, fallback_company=fallback_company, fallback_loc=loc)
                page2.close()
            except Exception:
                pass
            if not is_valid_linkedin_job(parsed["title"], parsed["company"], parsed["jd_text"]):
                continue
            jobs.append(
                {
                    "job_url": href,
                    "application_url": parsed.get("application_url", href),
                    "title": dedupe_repeated_phrase(parsed["title"]),
                    "company": parsed["company"],
                    "location": parsed["location"],
                    "job_posted_date": parsed.get("job_posted_date", ""),
                    "jd_text": parsed["jd_text"],
                    "notes": parsed["notes"],
                    "quick_apply": parsed["quick_apply"],
                }
            )
            if len(jobs) >= limit:
                break
        if not use_system_edge or created_context:
            context.close()
        if not use_system_edge and browser is not None:
            browser.close()
    return jobs[:limit]


def write_jobs_csv(rows: List[dict], platform: str, append: bool = False) -> None:
    fields = ["job_id", "platform", "company", "title", "location", "job_url", "application_url", "job_posted_date", "jd_text", "notes"]
    existing: List[dict] = []
    if append and JOBS_FILE.exists():
        with JOBS_FILE.open("r", encoding="utf-8-sig", newline="") as f:
            existing = list(csv.DictReader(f))
    out = existing[:]
    start_idx = len(existing) + 1
    for i, r in enumerate(rows, start=start_idx):
        out.append(
            {
                "job_id": r.get("job_id") or f"{platform}_scrape_{i:03d}",
                "platform": platform,
                "company": r.get("company", ""),
                "title": r.get("title", ""),
                "location": r.get("location", ""),
                "job_url": r.get("job_url", ""),
                "application_url": r.get("application_url", r.get("job_url", "")),
                "job_posted_date": normalize_posted_date(r.get("job_posted_date", "")),
                "jd_text": r.get("jd_text", ""),
                "notes": r.get("notes", ""),
            }
        )
    try:
        with JOBS_FILE.open("w", encoding="utf-8-sig", newline="") as f:
            w = csv.DictWriter(f, fieldnames=fields)
            w.writeheader()
            w.writerows(out)
    except PermissionError:
        fallback = JOBS_FILE.with_name("jobs_latest.csv")
        with fallback.open("w", encoding="utf-8-sig", newline="") as f:
            w = csv.DictWriter(f, fieldnames=fields)
            w.writeheader()
            w.writerows(out)
        print(f"Warning: jobs.csv is locked. Wrote to fallback file: {fallback}")


def _normalize_seek_text(value: str) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip().lower())


def _contains_any_term(text: str, terms: List[str]) -> bool:
    low = _normalize_seek_text(text)
    return any(term and _normalize_seek_text(term) in low for term in terms)


def _count_terms(text: str, terms: List[str]) -> int:
    low = _normalize_seek_text(text)
    return sum(1 for term in terms if term and _normalize_seek_text(term) in low)


def _seek_bucket_title_priority(title: str, strategy: dict) -> Tuple[int, str]:
    low = _normalize_seek_text(title)
    for bucket_name, score, key in [
        ("high", 35, "high_priority_titles"),
        ("medium", 22, "medium_priority_titles"),
        ("low", 10, "low_priority_titles"),
    ]:
        titles = [str(x) for x in strategy.get(key, []) if str(x).strip()]
        if any(_normalize_seek_text(term) in low for term in titles):
            return score, bucket_name
    return 0, "other"


def _seek_hard_mismatch_reason(text: str, strategy: dict) -> str:
    low = _normalize_seek_text(text)
    for phrase in [str(x) for x in strategy.get("hard_mismatch_phrases", []) if str(x).strip()]:
        norm_phrase = _normalize_seek_text(phrase)
        if norm_phrase and norm_phrase in low:
            return phrase
    return ""


def rank_seek_job(job: dict, strategy: dict) -> Tuple[float, List[str], str]:
    title = str(job.get("title", "")).strip()
    jd_text = str(job.get("jd_text", "")).strip()
    combined = f"{title}\n{jd_text}"
    low = _normalize_seek_text(combined)

    score = 0.0
    reasons: List[str] = []

    title_score, bucket = _seek_bucket_title_priority(title, strategy)
    score += title_score
    if title_score:
        reasons.append(f"title:{bucket}")

    keyword_groups = strategy.get("keyword_groups", {}) if isinstance(strategy.get("keyword_groups"), dict) else {}
    group_weights = {
        "data_reporting": 3.2,
        "accounting_finance": 3.0,
        "business_process": 3.0,
        "ai_fused_business": 2.4,
    }
    for group_name, weight in group_weights.items():
        terms = [str(x) for x in keyword_groups.get(group_name, []) if str(x).strip()]
        hits = _count_terms(combined, terms)
        if hits:
            score += min(hits, 6) * weight
            reasons.append(f"{group_name}:{hits}")

    if _contains_any_term(title, [str(x) for x in strategy.get("ai_fused_business_bucket", [])]):
        score += 10.0
        reasons.append("ai_fused_title")

    if _contains_any_term(combined, [str(x) for x in strategy.get("downrank_pure_technical_ai", [])]):
        score -= 60.0
        reasons.append("pure_technical_ai")

    exclusion_terms = [str(x) for x in strategy.get("exclusion_keywords", []) if str(x).strip()]
    exclusion_hits = _count_terms(combined, exclusion_terms)
    if exclusion_hits:
        score -= exclusion_hits * 8.0
        reasons.append(f"exclusion:{exclusion_hits}")

    mismatch = _seek_hard_mismatch_reason(combined, strategy)
    if mismatch:
        score -= 120.0
        reasons.append(f"hard_mismatch:{mismatch}")

    if _contains_any_term(combined, ["business analyst", "operations analyst", "process analyst", "business support officer"]) and _contains_any_term(combined, keyword_groups.get("business_process", [])):
        score += 12.0
        reasons.append("ba_process_fit")
    if _contains_any_term(combined, ["reporting analyst", "data analyst", "data visualisation analyst", "business intelligence analyst"]) and _contains_any_term(combined, keyword_groups.get("data_reporting", [])):
        score += 12.0
        reasons.append("data_reporting_fit")
    if _contains_any_term(combined, ["graduate accountant", "assistant accountant", "finance officer", "accounts officer"]) and _contains_any_term(combined, keyword_groups.get("accounting_finance", [])):
        score += 12.0
        reasons.append("accounting_fit")
    if _contains_any_term(combined, ["ai", "automation", "digital transformation", "innovation"]) and _contains_any_term(combined, keyword_groups.get("ai_fused_business", [])):
        score += 8.0
        reasons.append("ai_business_fit")

    return round(score, 1), reasons, mismatch


def build_seek_strategy_report(ranked_jobs: List[dict], strategy: dict, location: str) -> str:
    lines = [
        "# Seek Search Strategy",
        "",
        f"- Location: {location}",
        f"- Query templates: {len(strategy.get('query_templates', []))}",
        f"- Ranked jobs kept: {len(ranked_jobs)}",
        "",
        "## Query Templates",
        "",
    ]
    for item in strategy.get("query_templates", []):
        lines.append(f"- [{item.get('bucket','?')}] {item.get('label','query')}: {item.get('query','')}")
    lines.extend(["", "## Top Ranked Jobs", ""])
    for job in ranked_jobs[:20]:
        lines.append(
            f"- {job.get('search_score', 0)} | {job.get('company','')} | {job.get('title','')} | {job.get('job_posted_date','')} | {job.get('job_url','')} | {job.get('search_reasons','')}"
        )
    return "\n".join(lines) + "\n"


def scrape_seek_strategy(
    *,
    location: str,
    per_query_limit: int = 10,
    max_results: int = 40,
    append: bool = True,
    channel: str = "msedge",
    headless: bool = False,
    use_system_edge: bool = True,
    cdp_url: str = "http://127.0.0.1:9222",
) -> List[dict]:
    strategy = load_seek_search_strategy()
    deduped: Dict[str, dict] = {}
    for item in strategy.get("query_templates", []):
        query = str(item.get("query", "")).strip()
        if not query:
            continue
        rows = scrape_seek_jobs(
            query=query,
            location=location,
            limit=per_query_limit,
            channel=channel,
            headless=headless,
            use_system_edge=use_system_edge,
            cdp_url=cdp_url,
        )
        for row in rows:
            url = str(row.get("job_url", "")).strip()
            if not url:
                continue
            ranked_score, reasons, mismatch = rank_seek_job(row, strategy)
            row = dict(row)
            row["search_bucket"] = str(item.get("bucket", "")).strip()
            row["search_query_label"] = str(item.get("label", "")).strip()
            row["search_query"] = query
            row["search_score"] = ranked_score
            row["search_reasons"] = ", ".join(reasons)
            if mismatch:
                row["notes"] = f"{row.get('notes','')}; search_filtered:{mismatch}".strip("; ")
            existing = deduped.get(url)
            if existing is None or float(row.get("search_score", 0) or 0) > float(existing.get("search_score", 0) or 0):
                deduped[url] = row

    ranked_jobs = sorted(
        [row for row in deduped.values() if float(row.get("search_score", 0) or 0) > 0],
        key=lambda row: (float(row.get("search_score", 0) or 0), normalize_posted_date(row.get("job_posted_date", ""))),
        reverse=True,
    )
    ranked_jobs = ranked_jobs[:max_results]

    SEEK_STRATEGY_REPORT_FILE.parent.mkdir(parents=True, exist_ok=True)
    SEEK_STRATEGY_REPORT_FILE.write_text(build_seek_strategy_report(ranked_jobs, strategy, location), encoding="utf-8")
    write_jobs_csv(ranked_jobs, platform="seek", append=append)
    return ranked_jobs


def detect_portal_api(company: dict) -> Optional[dict]:
    api_value = str(company.get("api", "")).strip()
    if "greenhouse" in api_value:
        return {"type": "greenhouse", "url": api_value}
    careers_url = str(company.get("careers_url", "")).strip()
    ashby = re.search(r"jobs\.ashbyhq\.com/([^/?#]+)", careers_url, flags=re.I)
    if ashby:
        slug = ashby.group(1)
        return {
            "type": "ashby",
            "url": f"https://api.ashbyhq.com/posting-api/job-board/{slug}?includeCompensation=true",
        }
    lever = re.search(r"jobs\.lever\.co/([^/?#]+)", careers_url, flags=re.I)
    if lever:
        slug = lever.group(1)
        return {"type": "lever", "url": f"https://api.lever.co/v0/postings/{slug}"}
    gh = re.search(r"job-boards(?:\.eu)?\.greenhouse\.io/([^/?#]+)", careers_url, flags=re.I)
    if gh:
        slug = gh.group(1)
        return {"type": "greenhouse", "url": f"https://boards-api.greenhouse.io/v1/boards/{slug}/jobs"}
    return None


def parse_greenhouse_jobs(data: dict, company_name: str) -> List[dict]:
    jobs = data.get("jobs", []) if isinstance(data, dict) else []
    out = []
    for j in jobs:
        if not isinstance(j, dict):
            continue
        out.append(
            {
                "title": str(j.get("title", "")).strip(),
                "job_url": str(j.get("absolute_url", "")).strip(),
                "company": company_name,
                "location": str((j.get("location") or {}).get("name", "")).strip() if isinstance(j.get("location"), dict) else "",
                "job_posted_date": normalize_posted_date(j.get("updated_at", "") or j.get("first_published", "")),
                "jd_text": str(j.get("content", "")).strip(),
                "notes": "greenhouse_api",
            }
        )
    return out


def parse_ashby_jobs(data: dict, company_name: str) -> List[dict]:
    jobs = data.get("jobs", []) if isinstance(data, dict) else []
    out = []
    for j in jobs:
        if not isinstance(j, dict):
            continue
        out.append(
            {
                "title": str(j.get("title", "")).strip(),
                "job_url": str(j.get("jobUrl", "")).strip(),
                "company": company_name,
                "location": str(j.get("location", "")).strip(),
                "job_posted_date": normalize_posted_date(j.get("publishedDate", "") or j.get("createdAt", "")),
                "jd_text": str(j.get("descriptionHtml", "")).strip()[:12000],
                "notes": "ashby_api",
            }
        )
    return out


def parse_lever_jobs(data: list, company_name: str) -> List[dict]:
    out = []
    if not isinstance(data, list):
        return out
    for j in data:
        if not isinstance(j, dict):
            continue
        categories = j.get("categories")
        location = ""
        if isinstance(categories, dict):
            location = str(categories.get("location", "")).strip()
        posted_raw = j.get("createdAt", "")
        if isinstance(posted_raw, (int, float)):
            try:
                posted_raw = datetime.fromtimestamp(float(posted_raw) / 1000.0).date().isoformat()
            except Exception:
                posted_raw = ""
        out.append(
            {
                "title": str(j.get("text", "")).strip(),
                "job_url": str(j.get("hostedUrl", "")).strip(),
                "company": company_name,
                "location": location,
                "job_posted_date": normalize_posted_date(posted_raw),
                "jd_text": str(j.get("description", "")).strip()[:12000],
                "notes": "lever_api",
            }
        )
    return out


def build_title_filter(title_filter: dict):
    positive = [str(x).lower().strip() for x in (title_filter.get("positive", []) if isinstance(title_filter, dict) else [])]
    negative = [str(x).lower().strip() for x in (title_filter.get("negative", []) if isinstance(title_filter, dict) else [])]

    def _accept(title: str) -> bool:
        low = str(title or "").lower()
        has_pos = (not positive) or any(k and k in low for k in positive)
        has_neg = any(k and k in low for k in negative)
        return has_pos and not has_neg

    return _accept


def load_seen_urls_for_scan() -> set:
    seen = set()
    if SCAN_HISTORY_FILE.exists():
        with SCAN_HISTORY_FILE.open("r", encoding="utf-8", errors="ignore") as f:
            for i, line in enumerate(f):
                if i == 0:
                    continue
                parts = line.strip().split("\t")
                if parts and parts[0].strip():
                    seen.add(parts[0].strip().lower())
    if PIPELINE_FILE.exists():
        text = PIPELINE_FILE.read_text(encoding="utf-8", errors="ignore")
        for m in re.finditer(r"https?://\S+", text):
            seen.add(m.group(0).strip().rstrip("|").lower())
    if JOBS_FILE.exists():
        with JOBS_FILE.open("r", encoding="utf-8-sig", newline="") as f:
            for row in csv.DictReader(f):
                url = str(row.get("job_url", "")).strip().lower()
                if url:
                    seen.add(url)
    return seen


def ensure_pipeline_files() -> None:
    PIPELINE_FILE.parent.mkdir(parents=True, exist_ok=True)
    if not PIPELINE_FILE.exists():
        PIPELINE_FILE.write_text("# Pipeline Inbox\n\n## Pending\n\n## Processed\n", encoding="utf-8")
    if not SCAN_HISTORY_FILE.exists():
        SCAN_HISTORY_FILE.write_text("url\tfirst_seen\tportal\ttitle\tcompany\tstatus\n", encoding="utf-8")


def append_to_pipeline(offers: List[dict]) -> None:
    if not offers:
        return
    ensure_pipeline_files()
    text = PIPELINE_FILE.read_text(encoding="utf-8")
    marker = "## Pending"
    if marker not in text:
        text = text.rstrip() + f"\n\n{marker}\n"
    idx = text.find(marker)
    insert_after = text.find("\n", idx)
    if insert_after == -1:
        insert_after = len(text)
    block = "\n" + "\n".join(
        f"- [ ] {o.get('job_url','')} | {o.get('company','')} | {o.get('title','')}" for o in offers
    )
    new_text = text[: insert_after + 1] + block + text[insert_after + 1 :]
    PIPELINE_FILE.write_text(new_text, encoding="utf-8")


def append_scan_history(offers: List[dict], scan_date: str) -> None:
    if not offers:
        return
    ensure_pipeline_files()
    with SCAN_HISTORY_FILE.open("a", encoding="utf-8", newline="") as f:
        for o in offers:
            row = [
                str(o.get("job_url", "")).strip(),
                scan_date,
                str(o.get("notes", "")).strip(),
                str(o.get("title", "")).replace("\t", " ").strip(),
                str(o.get("company", "")).replace("\t", " ").strip(),
                "added",
            ]
            f.write("\t".join(row) + "\n")


def scan_portals(
    config_path: Path = PORTALS_FILE,
    company_filter: str = "",
    dry_run: bool = False,
    append_jobs: bool = True,
) -> None:
    if not config_path.exists():
        raise FileNotFoundError(f"Portals config not found: {config_path}. Run init first.")
    with config_path.open("r", encoding="utf-8") as f:
        cfg = yaml.safe_load(f) or {}
    companies = cfg.get("tracked_companies", [])
    if not isinstance(companies, list):
        companies = []
    title_accept = build_title_filter(cfg.get("title_filter", {}))
    filter_value = company_filter.strip().lower()
    targets = []
    for comp in companies:
        if not isinstance(comp, dict):
            continue
        if comp.get("enabled", True) is False:
            continue
        name = str(comp.get("name", "")).strip()
        if filter_value and filter_value not in name.lower():
            continue
        api = detect_portal_api(comp)
        if not api:
            continue
        targets.append({"name": name, "api": api})

    if not targets:
        print("No portal targets matched. Check config/portals.yaml and enabled flags.")
        return

    seen_urls = load_seen_urls_for_scan()
    parsed_offers: List[dict] = []
    parser_map = {
        "greenhouse": parse_greenhouse_jobs,
        "ashby": parse_ashby_jobs,
        "lever": parse_lever_jobs,
    }

    def _fetch_target(target: dict) -> List[dict]:
        name = target["name"]
        api = target["api"]
        parser = parser_map.get(api.get("type"))
        if not parser:
            return []
        try:
            with httpx.Client(timeout=15, follow_redirects=True) as client:
                resp = client.get(api.get("url", ""))
                resp.raise_for_status()
                payload = resp.json()
            rows = parser(payload, name)
            accepted = []
            for r in rows:
                if not title_accept(r.get("title", "")):
                    continue
                url = str(r.get("job_url", "")).strip().lower()
                if not url or url in seen_urls:
                    continue
                seen_urls.add(url)
                accepted.append(r)
            return accepted
        except Exception:
            return []

    with concurrent.futures.ThreadPoolExecutor(max_workers=min(10, len(targets))) as ex:
        futures = [ex.submit(_fetch_target, t) for t in targets]
        for fut in concurrent.futures.as_completed(futures):
            parsed_offers.extend(fut.result())

    scan_date = datetime.now().strftime("%Y-%m-%d")
    parsed_offers.sort(key=lambda x: (x.get("company", ""), x.get("title", "")))
    print(f"Portal scan date: {scan_date}")
    print(f"Targets scanned: {len(targets)}")
    print(f"New offers found: {len(parsed_offers)}")
    for o in parsed_offers[:25]:
        print(f"+ {o.get('company','')} | {o.get('title','')} | {o.get('location','')}")

    if dry_run or not parsed_offers:
        if dry_run:
            print("Dry run enabled: no files written.")
        return

    append_to_pipeline(parsed_offers)
    append_scan_history(parsed_offers, scan_date)
    if append_jobs:
        write_jobs_csv(parsed_offers, platform="portal", append=True)
    print(f"Pipeline updated: {PIPELINE_FILE}")
    print(f"Scan history updated: {SCAN_HISTORY_FILE}")
    if append_jobs:
        print(f"Jobs appended: {JOBS_FILE}")


def assist_apply_dialog(
    platform: str,
    limit: int = 20,
    headless: bool = False,
    auto_click: bool = False,
    auto_fill: bool = False,
    auto_submit: bool = False,
    channel: str = "msedge",
    use_system_edge: bool = False,
    cdp_url: str = "http://127.0.0.1:9222",
) -> None:
    ensure_playwright_ready()
    if not AUDIT_XLSX_FILE.exists():
        raise FileNotFoundError(f"Audit file not found: {AUDIT_XLSX_FILE}. Run pipeline first.")

    rows = read_xlsx_rows(AUDIT_XLSX_FILE)
    pending_pairs: List[Tuple[int, dict]] = []
    for idx, row in enumerate(rows):
        if row.get("platform", "").lower() == platform.lower() and row.get("status", "") == "ready_to_apply":
            pending_pairs.append((idx, row))
    pending_pairs = pending_pairs[:limit]

    if not pending_pairs:
        print(f"No pending rows for platform={platform}")
        return

    profile = load_candidate_profile()

    def detect_submission_success(page, platform_name: str) -> bool:
        try:
            current_url = page.url.lower()
        except Exception:
            current_url = ""
        try:
            body_text = (page.inner_text("body") or "").lower()
        except Exception:
            body_text = ""
        markers = [
            "application submitted",
            "thanks for applying",
            "thank you for applying",
            "you've applied",
            "your application has been submitted",
        ]
        if any(x in body_text for x in markers):
            return True
        if platform_name.lower() == "linkedin":
            return ("/jobs/view/" in current_url and "submittedapplication=true" in current_url) or ("application sent" in body_text)
        if platform_name.lower() == "seek":
            return ("seek.com.au" in current_url and ("application" in current_url and "submitted" in current_url))
        return False

    def advance_application_steps(page, platform_name: str, max_steps: int = 5) -> List[str]:
        steps: List[str] = []
        base_selectors = [
            "button:has-text('Continue')",
            "button:has-text('Next')",
            "button:has-text('Review')",
            "button:has-text('Review application')",
            "a:has-text('Continue')",
        ]
        if platform_name.lower() == "linkedin":
            base_selectors = [
                "button:has-text('Continue to next step')",
                "button:has-text('Next')",
                "button:has-text('Review')",
                "button:has-text('Review application')",
                "button[aria-label*='Continue']",
            ] + base_selectors
        if platform_name.lower() == "seek":
            base_selectors = [
                "button:has-text('Continue')",
                "button:has-text('Save and continue')",
                "a:has-text('Continue')",
            ] + base_selectors
        for _ in range(max_steps):
            moved = click_first(page, base_selectors)
            if not moved:
                break
            steps.append("advanced_step")
            page.wait_for_timeout(900)
        return steps

    with sync_playwright() as p:
        browser = None
        context = None
        if use_system_edge and not storage_state_path(platform.lower()).exists():
            ensure_system_edge_cdp(cdp_url, start_url="https://www.seek.com.au/" if platform.lower() == "seek" else "https://www.linkedin.com/")
            try:
                browser = p.chromium.connect_over_cdp(cdp_url)
            except Exception as e:
                print(
                    "Warning: failed to attach system Edge CDP; falling back to local saved browser profile. "
                    "For strict system Edge mode, fully close Edge then run: "
                    "msedge --remote-debugging-port=9222"
                )
                browser, context, _ = open_saved_platform_context(
                    p,
                    platform.lower(),
                    headless=headless,
                    channel=channel,
                )
                use_system_edge = False
            else:
                context = browser.contexts[0] if browser.contexts else browser.new_context()
        else:
            browser, context, _ = open_saved_platform_context(
                p,
                platform.lower(),
                headless=headless,
                channel=channel,
            )
            use_system_edge = False
        page = context.new_page()
        for idx, row in pending_pairs:
            url = str(row.get("job_url", "")).strip()
            if not url:
                rows[idx]["status"] = "skipped_no_url"
                continue

            print("=" * 60)
            print(f"[Dialog] {row.get('title', '')} @ {row.get('company', '')}")
            print(f"URL: {url}")
            print(f"Resume: {row.get('resume_file', '')}")
            print(f"Cover Letter: {row.get('cover_letter_file', '')}")
            try:
                page.goto(url, wait_until="domcontentloaded")
            except PlaywrightTimeoutError:
                pass

            auto_msg = ""
            if auto_click:
                chain = click_apply_chain(page)
                if not chain:
                    auto_msg = try_auto_click_apply(page, platform)
                else:
                    auto_msg = " -> ".join(chain)
                print(f"Auto step: {auto_msg}")

            fill_stats = {"fields": 0, "uploads": 0, "coverletter": 0}
            if auto_fill:
                cover_text = read_cover_letter_text(str(row.get("cover_letter_file", "")))
                fill_stats = autofill_form(
                    page,
                    profile=profile,
                    resume_file=str(row.get("resume_file", "")),
                    cover_letter_text=cover_text,
                )
                auto_msg = f"{auto_msg}; autofill fields={fill_stats['fields']} upload={fill_stats['uploads']} cover={fill_stats['coverletter']}".strip("; ")
                print(f"Autofill: {auto_msg}")

            submitted = False
            if auto_submit:
                advance_steps = advance_application_steps(page, platform_name=platform, max_steps=5)
                if advance_steps:
                    auto_msg = f"{auto_msg}; {' -> '.join(advance_steps)}".strip("; ")
                submitted = click_first(
                    page,
                    [
                        "button:has-text('Submit')",
                        "button:has-text('Submit application')",
                        "button:has-text('Submit Application')",
                        "button:has-text('Send application')",
                        "button:has-text('Apply now')",
                        "button:has-text('Done')",
                        "input[type='submit']",
                    ],
                )
                if not submitted:
                    submitted = detect_submission_success(page, platform)
                if submitted:
                    auto_msg = f"{auto_msg}; submitted"
                    print("Auto submit: submitted")
                else:
                    auto_msg = f"{auto_msg}; submit_not_found"
                    print("Auto submit: submit button not found")
            # Safety lock: if we couldn't upload resume, never mark submitted automatically.
            if auto_fill and "upload=0" in auto_msg and submitted:
                submitted = False
                auto_msg = f"{auto_msg}; blocked_no_resume_upload"
                print("Safety lock: blocked submit because resume upload was not detected.")

            if submitted:
                decision = "submitted"
            elif auto_fill or auto_click:
                decision = "ready_for_review"
            else:
                decision = "hold"
            rows[idx]["status"] = decision
            rows[idx]["auto_step"] = auto_msg
            rows[idx]["applied_at"] = datetime.now().isoformat(timespec="seconds")
        if not use_system_edge:
            context.close()
        if browser is not None:
            browser.close()

    sync_tracker_outputs(rows)
    print(f"Audit updated: {AUDIT_XLSX_FILE}")


def tracker_report() -> None:
    rows = read_xlsx_rows(AUDIT_XLSX_FILE)
    final_rows = sync_tracker_outputs(rows)
    summary = summarize_tracker(final_rows)
    print("=" * 60)
    print("Application Tracker Summary")
    print("=" * 60)
    print(f"Total roles: {summary['total']}")
    print(f"Average match score: {summary['avg_match_score']}")
    print(f"Resume coverage: {summary['with_resume_pct']}%")
    print(f"Cover letter coverage: {summary['with_cover_pct']}%")
    print(f"High-match rate (>=70): {summary.get('high_match_rate_pct', 0.0)}%")
    print(f"Submission rate: {summary.get('submission_rate_pct', 0.0)}%")
    print(f"Interview-or-better rate: {summary.get('interview_plus_rate_pct', 0.0)}%")
    print(f"Fresh jobs (posted <=7 days): {summary.get('fresh_jobs_7d_pct', 0.0)}%")
    print(f"Aged jobs (posted >14 days): {summary.get('aged_jobs_14d_pct', 0.0)}%")
    print(f"Posted-date coverage: {summary.get('posted_date_coverage_pct', 0.0)}%")
    print(f"Avg days since posted: {summary.get('avg_days_since_posted', 0.0)}")
    for status, count in summary["counts"].items():
        print(f"- {status}: {count}")
    print(f"CSV tracker: {TRACKER_CSV_FILE}")
    print(f"Markdown tracker: {TRACKER_MD_FILE}")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Resume tailoring + application workflow")
    sub = parser.add_subparsers(dest="cmd", required=True)

    sub.add_parser("init", help="Create folders + starter files")

    auth = sub.add_parser("auth", help="Open browser for manual login and save session")
    auth.add_argument("--platform", required=True, choices=["seek", "linkedin"])
    auth.add_argument("--headless", action="store_true")
    auth.add_argument("--wait-seconds", type=int, default=120)
    auth.add_argument("--channel", choices=["msedge", "chrome"], default="msedge")
    auth.add_argument("--use-system-edge", action=argparse.BooleanOptionalAction, default=True)
    auth.add_argument("--cdp-port", type=int, default=9222)

    run = sub.add_parser("run", help="Build tailored resumes + score + report")
    run.add_argument("--use-openai", action="store_true", help="Use OPENAI_API_KEY for rewrite")

    sub.add_parser("tracker", help="Normalize, dedupe, and export tracker views")
    hydrate_cmd = sub.add_parser("hydrate-profile", help="Import supplemental profile facts from external profile project")
    hydrate_cmd.add_argument(
        "--source-dir",
        default=r"D:\上课\Ai agent\digital twin",
        help="External profile project directory",
    )
    hydrate_cmd.add_argument("--max-facts", type=int, default=80, help="Maximum facts to keep")
    scan_portals_cmd = sub.add_parser("scan-portals", help="Scan Greenhouse/Ashby/Lever portals via public APIs")
    scan_portals_cmd.add_argument("--config", default=str(PORTALS_FILE), help="Portal config yaml path")
    scan_portals_cmd.add_argument("--company", default="", help="Scan only matching company name")
    scan_portals_cmd.add_argument("--dry-run", action="store_true", help="Preview without writing files")
    scan_portals_cmd.add_argument("--no-append-jobs", action="store_true", help="Do not append scan results into jobs/jobs.csv")

    scrape_cmd = sub.add_parser("scrape", help="Scrape jobs only (no application)")
    scrape_cmd.add_argument("--platform", required=True, choices=["seek", "linkedin"])
    scrape_cmd.add_argument("--query", default="business graduate")
    scrape_cmd.add_argument("--location", default="Brisbane QLD")
    scrape_cmd.add_argument("--limit", type=int, default=20)
    scrape_cmd.add_argument("--append", action="store_true", help="Append to existing jobs.csv")
    scrape_cmd.add_argument("--channel", choices=["msedge", "chrome"], default="msedge")
    scrape_cmd.add_argument("--headless", action="store_true")
    scrape_cmd.add_argument("--use-system-edge", action=argparse.BooleanOptionalAction, default=True)
    scrape_cmd.add_argument("--cdp-url", default="http://127.0.0.1:9222")

    seek_strategy_cmd = sub.add_parser("scrape-seek-strategy", help="Run bucketed Seek search strategy with ranking and exclusions")
    seek_strategy_cmd.add_argument("--location", default="Brisbane QLD")
    seek_strategy_cmd.add_argument("--per-query-limit", type=int, default=10)
    seek_strategy_cmd.add_argument("--max-results", type=int, default=40)
    seek_strategy_cmd.add_argument("--append", action="store_true", help="Append ranked results into jobs.csv")
    seek_strategy_cmd.add_argument("--channel", choices=["msedge", "chrome"], default="msedge")
    seek_strategy_cmd.add_argument("--headless", action="store_true")
    seek_strategy_cmd.add_argument("--use-system-edge", action=argparse.BooleanOptionalAction, default=True)
    seek_strategy_cmd.add_argument("--cdp-url", default="http://127.0.0.1:9222")

    apply_cmd = sub.add_parser("apply", help="Assist manual applying using saved session")
    apply_cmd.add_argument("--platform", required=True, choices=["seek", "linkedin"])
    apply_cmd.add_argument("--limit", type=int, default=20)
    apply_cmd.add_argument("--headless", action="store_true")
    apply_cmd.add_argument("--auto-click", action="store_true", help="Try clicking Apply/Easy Apply automatically")
    apply_cmd.add_argument("--auto-fill", action="store_true", help="Auto-fill detected application fields")
    apply_cmd.add_argument("--auto-submit", action="store_true", help="Attempt final submit automatically")
    apply_cmd.add_argument("--channel", choices=["msedge", "chrome"], default="msedge")
    apply_cmd.add_argument("--use-system-edge", action=argparse.BooleanOptionalAction, default=True)
    apply_cmd.add_argument("--cdp-url", default="http://127.0.0.1:9222")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    if args.cmd == "init":
        init_files()
        print("Initialized project scaffolding.")
    elif args.cmd == "auth":
        login_platform(
            args.platform,
            headless=args.headless,
            wait_seconds=args.wait_seconds,
            channel=args.channel,
            use_system_edge=args.use_system_edge,
            cdp_port=args.cdp_port,
        )
    elif args.cmd == "run":
        run_pipeline(use_openai=args.use_openai)
    elif args.cmd == "tracker":
        tracker_report()
    elif args.cmd == "hydrate-profile":
        hydrate_profile_from_external(source_dir=args.source_dir, max_facts=args.max_facts)
    elif args.cmd == "scan-portals":
        scan_portals(
            config_path=Path(args.config),
            company_filter=args.company,
            dry_run=args.dry_run,
            append_jobs=not args.no_append_jobs,
        )
    elif args.cmd == "scrape-seek-strategy":
        rows = scrape_seek_strategy(
            location=args.location,
            per_query_limit=args.per_query_limit,
            max_results=args.max_results,
            append=args.append,
            channel=args.channel,
            headless=args.headless,
            use_system_edge=args.use_system_edge,
            cdp_url=args.cdp_url,
        )
        print(f"Ranked {len(rows)} Seek jobs to: {JOBS_FILE}")
        print(f"Strategy report: {SEEK_STRATEGY_REPORT_FILE}")
    elif args.cmd == "scrape":
        if args.platform == "seek":
            rows = scrape_seek_jobs(
                args.query,
                args.location,
                args.limit,
                channel=args.channel,
                headless=args.headless,
                use_system_edge=args.use_system_edge,
                cdp_url=args.cdp_url,
            )
        else:
            rows = scrape_linkedin_jobs(
                args.query,
                args.location,
                args.limit,
                channel=args.channel,
                headless=args.headless,
                use_system_edge=args.use_system_edge,
                cdp_url=args.cdp_url,
            )
        write_jobs_csv(rows, platform=args.platform, append=args.append)
        print(f"Scraped {len(rows)} jobs to: {JOBS_FILE}")
    elif args.cmd == "apply":
        assist_apply_dialog(
            platform=args.platform,
            limit=args.limit,
            headless=args.headless,
            auto_click=args.auto_click,
            auto_fill=args.auto_fill,
            auto_submit=args.auto_submit,
            channel=args.channel,
            use_system_edge=args.use_system_edge,
            cdp_url=args.cdp_url,
        )


if __name__ == "__main__":
    main()
